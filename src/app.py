import os
import queue
import struct
import re
import sys
import threading
import webbrowser
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from getpass import getuser
import customtkinter as ctk
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import demo_config
import demo_index
VERSION = 'BOM Search Demo'
CHANGELOG = [('Portfolio demo', ['Search synthetic BOM metadata and spreadsheet contents.', 'Browse component details and see where a part is used.', 'Explore a sample revision-impact tree and export a checklist.'])]
ctk.set_default_color_theme('blue')
FONT = 'Segoe UI'
(TICKED, UNTICKED) = ('☑', '☐')
TABLE_COLOURS = {'light': dict(bg='#ffffff', alt='#f4f6f8', fg='#11181c', heading_bg='#e8ebef', heading_fg='#41505e', selected='#cfe3fb', selected_fg='#0b2545', border='#d5dbe1', prototype='#fbe0ea'), 'dark': dict(bg='#1f2429', alt='#252b31', fg='#e6eaee', heading_bg='#2b3238', heading_fg='#9fb0bf', selected='#2f4f76', selected_fg='#ffffff', border='#39424a', prototype='#5c2a40')}
PROTOTYPE_EXPORT_FILL = TABLE_COLOURS['light']['prototype'].lstrip('#').upper()
RESULT_COLUMNS = [('BOM #', 'BOM #', 150, 'center'), ('Drawing', 'Drawing', 160, 'center'), ('Assembly', 'Assembly Name', 240, 'center'), ('Customer', 'Customer / Platform', 175, 'center'), ('CustomerNo', 'Customer part no.', 150, 'center'), ('ServiceNo', 'Service part no.', 150, 'center'), ('Commodity', 'Commodity', 175, 'center'), ('Plant', 'Plant', 90, 'center'), ('Changed', 'Date', 105, 'center'), ('SignedBy', 'Signed by', 175, 'center'), ('MatchedIn', 'Matched in', 195, 'center'), ('MatchedValue', 'What matched', 400, 'center')]
COMPONENT_RESULT_COLUMNS = [('PartNumber', 'Part Number', 150, 'center'), ('PartName', 'Part Name', 260, 'center'), ('Drawing', 'Drawing', 160, 'center'), ('MaterialSpec', 'Material Spec', 200, 'center'), ('ThkMin', 'Thk min', 130, 'center'), ('ThkMax', 'Thk max', 130, 'center'), ('Finish', 'Surface Finish', 170, 'center'), ('Commodity', 'Commodity', 170, 'center'), ('UsedIn', 'Used in', 75, 'center'), ('MatchedIn', 'Matched in', 120, 'center')]
ROW_BUTTON_BORDER = ('white', 'white')
ROW_BUTTON_OFF = ('gray78', 'gray30')
ACTION_HELP = {'Open BOM': 'Open a synthetic spreadsheet in the demo dataset.', 'Drawing': 'Open a synthetic drawing PDF.', 'Customer Dwg': 'Open a synthetic customer-print example where available.', 'DFMEA': 'Open a synthetic risk-analysis record where available.', 'CAD': 'Open a synthetic CAD text file.', 'Components': 'List the parts inside the selected assembly.', 'Projects': 'Browse the synthetic project records linked to this assembly.'}
PRODUCTION_ONLY_ACTIONS = ('Customer Dwg', 'DFMEA', 'Projects')
for _label in PRODUCTION_ONLY_ACTIONS:
    if _label in ACTION_HELP:
        ACTION_HELP[_label] += ' Off on a prototype row: this is for released BOMs.'
ONE_ROW_ACTIONS = ('Components', 'Projects')
for _label in ONE_ROW_ACTIONS:
    ACTION_HELP[_label] += ' One row only.'
ROW_WHAT = {'Open BOM': 'BOM', 'Drawing': 'drawing', 'Customer Dwg': 'customer drawing', 'DFMEA': 'DFMEA', 'CAD': 'CAD file'}
PROTOTYPE_CHECKLIST_NOTE = 'The revision checklist is for released BOMs. This part is from a prototype BOM, which is not something an ECN reissues.\n\nSearch the part number to find the released BOMs that use it, and ask from one of those.'
DFMEA_NOT_BUILT = ('Not included in demo', 'Company-specific risk-analysis records are omitted from this synthetic portfolio dataset.')
DFMEA_NONE = ('No sample record', 'No synthetic risk-analysis record is linked to this assembly.')
DFMEA_NOT_NAMED = 'Risk-analysis records are omitted from this demo, so no such documents are named here.'
COLUMN_HELP = {'BOM #': 'The identifier assigned to this demo workbook.', 'Drawing': 'The drawing identifier associated with this assembly.', 'Assembly': 'The assembly name from the workbook.', 'Customer': 'A fictional customer and platform label.', 'CustomerNo': 'A fictional customer part number.', 'ServiceNo': 'A fictional service part number.', 'Commodity': 'The sample product group containing this workbook.', 'SignedBy': 'A synthetic reviewer name.', 'Changed': 'A synthetic date from the sample workbook.', 'MatchedIn': 'The field where the search term matched.', 'MatchedValue': 'The text that matched your search.', 'PartNumber': 'The synthetic part identifier.', 'PartName': 'The part name from a sample workbook.', 'MaterialSpec': 'The sample material specification.', 'ThkMin': 'The minimum sample material thickness.', 'ThkMax': 'The maximum sample material thickness.', 'Finish': 'The sample surface finish.', 'Plant': 'A fictional manufacturing location.'}
MODE_COLUMNS = {'boms': RESULT_COLUMNS, 'components': COMPONENT_RESULT_COLUMNS}
MODE_LABELS = {'boms': 'BOMs', 'components': 'Components'}
ROW_SOURCES = {'boms': {'BOM #': 'BOM #', 'Drawing': 'Drawing', 'Assembly': 'Assembly Name:', 'Customer': 'Customer & Platform:', 'CustomerNo': 'Customer Assembly No:', 'ServiceNo': 'Customer Service No:', 'Commodity': 'Folder', 'MatchedIn': 'MatchedIn', 'MatchedValue': 'MatchedValue'}, 'components': {'PartNumber': 'part_number', 'PartName': 'part_name', 'Drawing': 'drawing', 'MaterialSpec': 'material', 'ThkMin': 'thickness_min', 'ThkMax': 'thickness_max', 'Finish': 'finish', 'Commodity': 'Folder', 'UsedIn': 'BOMs', 'MatchedIn': 'MatchedIn'}}
MODE_ACCENTS = {'boms': {'normal': ('#1f6aa5', '#1f6aa5'), 'hover': ('#144870', '#144870'), 'panel': ('#e8eef5', '#20262c'), 'table': {'light': '#cfe3fb', 'dark': '#2f4f76'}}, 'components': {'normal': ('#2f8f4e', '#2f8f4e'), 'hover': ('#246b3b', '#246b3b'), 'panel': ('#e7f2ea', '#1e2622'), 'table': {'light': '#cbe9d4', 'dark': '#2c5138'}}}

def _is_prototype(row):
    return str(row.get('Source', '')) == demo_index.PROTOTYPE

def needs_prototype_rule(number):
    if demo_index.prototype_key(number) is not None:
        return True
    digits = re.match('\\d+', str(number).strip())
    return bool(digits) and len(digits.group(0)) < demo_index.DRAWING_NUMBER_LENGTH

def commodity_counts(results):
    if 'Source' in results.columns:
        results = results[results['Source'] != demo_index.PROTOTYPE]
    return results['Folder'].astype(str).value_counts()

def folder_label(row):
    folder = row.get('Folder', '')
    if _is_prototype(row):
        return f'Project {folder}'
    return demo_config.commodity_name(folder)

def prototype_note(index):
    if getattr(index, 'prototype', None) is not None and getattr(index, 'prototype_error', ''):
        return ' Prototype search failed - see the self-check.'
    if index.prototype_state() != 'ready':
        if not demo_config.paths().get('proto_bom_root'):
            return ''
        return ' Prototype folders not indexed yet.'
    when = index.prototype.built_on()
    return f' Prototypes: {index.prototype.workbook_count:,} BOMs' + (f', built {when}.' if when else '.')
NO_PROTOTYPE_ROOT = 'The prototype folders are not set up in this version yet, so this prototype file cannot be opened from here.'
PROTOTYPE_NOT_INDEXED = "The prototype folders have not been indexed yet, so this prototype part's files cannot be found from here."
PROTOTYPE_TAB = demo_index.PROTOTYPE

def _row_is_shown(row, chosen):
    if chosen == 'All':
        return True
    if chosen == PROTOTYPE_TAB:
        return _is_prototype(row)
    return str(row.get('Folder', '')) == chosen

def _one_line(value, limit=None):
    if value is None or value != value:
        return ''
    text = ' · '.join((part.strip() for part in str(value).splitlines() if part.strip()))
    if limit and len(text) > limit:
        text = text[:limit].rstrip() + '…'
    return text
MM_PER_INCH = Decimal('25.4')

def _thickness(value):
    if value is None or value != value:
        return ''
    text = str(value).strip()
    if not text:
        return ''
    try:
        inches = Decimal(text)
    except (ArithmeticError, ValueError):
        return text
    if not inches.is_finite():
        return text
    try:
        shown = str(inches.quantize(Decimal('0.001'), rounding=ROUND_HALF_UP))
        millimetres = (inches * MM_PER_INCH).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except ArithmeticError:
        return text
    if '.' in shown:
        shown = shown.rstrip('0').rstrip('.') or '0'
    return f'{shown} ({millimetres} mm)'

def _sort_rows(rows, descending):
    rows.sort(key=lambda pair: _sort_key(pair[0]), reverse=descending)
    rows.sort(key=lambda pair: _sort_key(pair[0])[0])
    return rows

def _sort_key(text):
    text = str(text).replace(',', '')
    try:
        return (0, float(text), '')
    except ValueError:
        pass
    try:
        return (0, float(text.split(' ', 1)[0]), '')
    except ValueError:
        return (1, 0.0, text.lower())

def settings_path():
    root = os.environ.get('LOCALAPPDATA', os.path.expanduser('~'))
    return os.path.join(root, 'BOM Search Demo', 'settings3.json')

def load_settings():
    import json
    try:
        with open(settings_path()) as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return {}

def save_settings(data):
    import json
    try:
        os.makedirs(os.path.dirname(settings_path()), exist_ok=True)
        with open(settings_path(), 'w') as handle:
            json.dump(data, handle)
    except OSError:
        pass

def stale_index_age(index):
    if not demo_config.warn_when_index_is_old():
        return None
    days = index.age_in_days()
    if days is None or days < demo_config.STALE_INDEX_DAYS:
        return None
    return days

def log_activity(action):
    try:
        path = demo_config.usage_log()
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, 'a', encoding='utf-8') as handle:
            handle.write(f'{getuser()},{datetime.now():%Y-%m-%d %H:%M:%S},{action}\n')
    except OSError:
        pass

def network_state(paths):
    cells = paths['cells']
    if not demo_config.network_ready(cells):
        return 'offline'
    (found, _) = demo_config.resolve_table(cells)
    return 'ready' if found else 'no index'

def copy_files_to_clipboard(paths, window):
    import ctypes
    from ctypes import wintypes
    (kernel32, user32) = (ctypes.windll.kernel32, ctypes.windll.user32)
    kernel32.GlobalAlloc.argtypes = [wintypes.UINT, ctypes.c_size_t]
    kernel32.GlobalAlloc.restype = wintypes.HGLOBAL
    kernel32.GlobalLock.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalUnlock.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalFree.argtypes = [wintypes.HGLOBAL]
    user32.OpenClipboard.argtypes = [wintypes.HWND]
    user32.SetClipboardData.argtypes = [wintypes.UINT, wintypes.HANDLE]
    user32.SetClipboardData.restype = wintypes.HANDLE
    names = ''.join((os.path.abspath(path) + '\x00' for path in paths)) + '\x00'
    payload = struct.pack('<IiiII', 20, 0, 0, 0, 1) + names.encode('utf-16-le')
    handle = kernel32.GlobalAlloc(66, len(payload))
    if not handle:
        raise OSError('Windows would not give the clipboard any memory.')
    pointer = kernel32.GlobalLock(handle)
    ctypes.memmove(pointer, payload, len(payload))
    kernel32.GlobalUnlock(handle)
    if not user32.OpenClipboard(window):
        kernel32.GlobalFree(handle)
        raise OSError('Another program is using the clipboard. Try again.')
    try:
        user32.EmptyClipboard()
        if not user32.SetClipboardData(15, handle):
            kernel32.GlobalFree(handle)
            raise OSError('Windows would not take the file onto the clipboard.')
    finally:
        user32.CloseClipboard()

def folder_files(folder, extensions):
    return [name for name in os.listdir(folder) if not os.path.isdir(os.path.join(folder, name)) and (extensions is None or name.lower().endswith(extensions))]
MANY_FILES = 10

class Batch:

    def __init__(self, owner, copying):
        (self.owner, self.copying) = (owner, copying)
        self.item = None
        self.files = {}
        self.missing = []
        self.notes = []
        self._listings = {}

    def collect(self, path, description):
        self.files.setdefault(path, set()).add(self.item)

    def listing(self, folder, extensions):
        key = (folder, extensions)
        if key not in self._listings:
            try:
                self._listings[key] = folder_files(folder, extensions)
            except OSError as error:
                self._listings[key] = error
        found = self._listings[key]
        if isinstance(found, OSError):
            raise found
        return found

    def note(self, text):
        if text not in self.notes:
            self.notes.append(text)

def batch_summary(copying, what, rows, found, files, missing, notes):
    verb = 'Copied' if copying else 'Opened'
    plural = what + 's'
    if not files:
        return ('\n\n'.join([f'No {what} for any of the {rows} selected rows.'] + notes), True)
    if missing or notes:
        text = f'{verb} {found} of {rows} {plural}.'
        if missing:
            listed = ', '.join(missing[:15])
            if len(missing) > 15:
                listed += f' ...and {len(missing) - 15} more'
            text += f' No {what} for: {listed}.'
        return ('\n\n'.join([text] + notes), True)
    if copying:
        return (f'{verb} {len(files)} {plural} - paste them into a folder or an email.', False)
    return (f'{verb} {len(files)} {plural}.', False)

def missing_file_message(path):
    if not demo_config.network_ready(path):
        return 'BOM Search Demo cannot reach the demo data folder.\n\nIf you are working from home, connect to the VPN and try again.\n\n' + path
    return f'This file is not there:\n\n{path}'

def dfmea_label(record):
    if not record:
        return ''
    if record['revision'] and record['next']:
        return f"DFMEA ({record['revision']} → {record['next']})"
    return 'DFMEA (see file)'

def update_text(owns_bom, covered, owner, dfmea):
    if owns_bom:
        text = f'BOM  (same drawing + CAD as {owner})' if covered else 'BOM + drawing + CAD'
        label = dfmea_label(dfmea)
        return f'{text} + {label}' if label else text
    return f'same drawing + CAD as {owner}' if covered else 'drawing + CAD'

def dfmea_actions(entries, lookup, root):
    rows = {}
    for entry in entries:
        if not entry.get('bom'):
            continue
        record = lookup(entry['part'])
        if not record:
            continue
        key = (record['folder'], record['filename'])
        if key in rows:
            held = rows[key][0]
            held['part'] += ', ' + entry['part']
            held['done'] = held['done'] and entry.get('done', False)
            continue
        what = f"Revise DFMEA  {record['revision']} → {record['next']}" if record['revision'] and record['next'] else 'Revise DFMEA  (revision: see file)'
        rows[key] = (dict(entry, done=entry.get('done', False)), what, record['filename'], os.path.join(root, record['folder']))
    return list(rows.values())

def bom_change_followups(state, count):
    followups = []
    if state != 'ready':
        followups.append(('Update the DFMEA', True, f"{count} BOM{('s are' if count != 1 else ' is')} being reissued. The DFMEA list has not been built, so DFMEAs are not named here - check each BOM's DFMEA by hand."))
    return followups + [('Check the packaging documents', False, 'Returnable, expendable, service - does this change reach any of them?'), ('Check the DVP', False, 'Does the change need validation testing, new or repeated?')]

def default_save_folder():
    return os.path.join(os.path.expanduser('~'), 'Documents')

def network_status_text(state):
    return 'BOM Search Demo cannot reach the network.' if state == 'offline' else 'The search index is not on the demo folder.'

def tooltip_position(pointer, size, window, gap=16, margin=8):
    ((px, py), (width, height), (left, top, right, bottom)) = (pointer, size, window)
    x = max(left + margin, min(px + gap, right - width - margin))
    y = py + gap
    if y + height > bottom - margin:
        y = py - gap - height
    return (x, max(top + margin, y))

class Tooltip:
    DELAY_MS = 450
    WRAP_PX = 320

    def __init__(self, widget, text):
        self.widget = widget
        self.text = (text or '').strip()
        self.window = None
        self.pending = None
        if not self.text:
            return
        for target in self._bindable(widget):
            target.bind('<Enter>', self._schedule, add='+')
            target.bind('<Leave>', self._hide, add='+')
            target.bind('<ButtonPress>', self._hide, add='+')

    @staticmethod
    def _bindable(widget):
        try:
            widget.bind('<Enter>', lambda event: None, add='+')
            return [widget]
        except (NotImplementedError, AttributeError, tk.TclError):
            pass
        inner = []
        for child in widget.winfo_children():
            inner.extend(Tooltip._bindable(child))
        return inner

    def _schedule(self, event=None):
        self._cancel()
        try:
            self.pending = self.widget.after(self.DELAY_MS, self._show)
        except tk.TclError:
            self.pending = None

    def _cancel(self):
        if self.pending is not None:
            try:
                self.widget.after_cancel(self.pending)
            except tk.TclError:
                pass
            self.pending = None

    def _show(self):
        self.pending = None
        if self.window is not None or not self.text:
            return
        try:
            pointer = self.widget.winfo_pointerxy()
            top = self.widget.winfo_toplevel()
            window = (top.winfo_rootx(), top.winfo_rooty(), top.winfo_rootx() + top.winfo_width(), top.winfo_rooty() + top.winfo_height())
        except tk.TclError:
            return
        dark = ctk.get_appearance_mode() == 'Dark'
        self.window = tk.Toplevel(self.widget)
        self.window.wm_overrideredirect(True)
        self.window.attributes('-topmost', True)
        tk.Label(self.window, text=self.text, justify='left', wraplength=self.WRAP_PX, background='#22262b' if dark else '#fbfbfc', foreground='#e6e8ea' if dark else '#1d2125', relief='solid', borderwidth=1, font=(FONT, 10), padx=10, pady=7).pack()
        self.window.update_idletasks()
        (x, y) = tooltip_position(pointer, (self.window.winfo_width(), self.window.winfo_height()), window)
        self.window.wm_geometry('+%d+%d' % (x, y))

    def _hide(self, event=None):
        self._cancel()
        if self.window is not None:
            try:
                self.window.destroy()
            except tk.TclError:
                pass
            self.window = None

def explain(widget, text):
    Tooltip(widget, text)
    return widget

class HeadingTooltip:

    def __init__(self, tree, texts):
        self.tree = tree
        self.texts = texts
        self.tip = Tooltip(tree, '')
        self.showing = None
        tree.bind('<Motion>', self._moved, add='+')
        tree.bind('<Leave>', lambda event: self._clear(), add='+')

    def _moved(self, event):
        if self.tree.identify_region(event.x, event.y) != 'heading':
            self._clear()
            return
        column = self.tree.identify_column(event.x)
        try:
            names = list(self.tree['columns'])
            key = names[int(column.replace('#', '')) - 1] if column != '#0' else '#0'
        except (ValueError, IndexError):
            self._clear()
            return
        if key == self.showing:
            return
        self._clear()
        text = self.texts.get(key)
        if not text:
            return
        self.showing = key
        self.tip.text = text
        self.tip._schedule()

    def _clear(self):
        self.showing = None
        self.tip._hide()

class App(ctk.CTk):

    def __init__(self):
        super().__init__()
        self.settings = load_settings()
        self.columns_panel = None
        ctk.set_appearance_mode(self.settings.get('theme', 'dark'))
        self.index = None
        self.index_error = None
        self.results = None
        self.projects_panel = None
        self.revision_panel = None
        self.network_card = None
        self._update_offered = None
        self._update_timer = None
        self.ui_queue = queue.Queue()
        self.title(VERSION if demo_config.is_live() else f'{VERSION}   —   {demo_config.describe_v2()}')
        icon = demo_config.app_icon()
        if icon:
            try:
                self.iconbitmap(default=icon)
            except tk.TclError:
                pass
        self.geometry('%dx%d' % (max(1100, int(self.winfo_screenwidth() * 0.94)), max(640, int(self.winfo_screenheight() * 0.9))))
        self.minsize(1100, 640)
        self.after(0, lambda : self.state('zoomed'))
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(4, weight=1)
        self._build_header()
        self._build_update_banner()
        self._build_stale_banner()
        self._build_search()
        self._build_results()
        self._build_components()
        self._build_status()
        self.bind('<Control-f>', lambda event: self._focus_search())
        self.bind('<Escape>', self._on_escape)
        self.protocol('WM_DELETE_WINDOW', self._on_close)
        self._apply_accent()
        self._drain_ui_queue()
        log_activity(f'App opened ({VERSION}, {demo_config.active_name()} profile)')
        threading.Thread(target=self._startup, daemon=True).start()

    def _build_header(self):
        header = ctk.CTkFrame(self, corner_radius=0, height=64)
        header.grid(row=0, column=0, sticky='ew')
        header.grid_columnconfigure(1, weight=1)
        header.grid_propagate(False)
        title = ctk.CTkLabel(header, text='BOM Search Demo', font=ctk.CTkFont(family=FONT, size=22, weight='bold'))
        title.grid(row=0, column=0, padx=(22, 10), pady=16, sticky='w')
        self.subtitle = ctk.CTkLabel(header, text='Synthetic engineering demo dataset', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray45', 'gray60'))
        self.subtitle.grid(row=0, column=1, padx=0, pady=(22, 14), sticky='w')
        self.theme_switch = ctk.CTkSegmentedButton(header, values=['Light', 'Dark'], width=150, command=self._on_theme_change, font=ctk.CTkFont(family=FONT, size=12))
        self.theme_switch.set('Dark' if ctk.get_appearance_mode() == 'Dark' else 'Light')
        self.theme_switch.grid(row=0, column=2, padx=(10, 12), pady=14)
        feedback = ctk.CTkButton(header, text='Feedback', width=90, command=self.send_feedback, fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=12))
        feedback.grid(row=0, column=3, padx=(0, 8), pady=14)
        explain(feedback, 'Send the maintainer a note about the app. Say what you searched for and what you expected to see - that is the part that makes a report fixable.')
        help_button = ctk.CTkButton(header, text='Help', width=70, command=self._show_help, font=ctk.CTkFont(family=FONT, size=12))
        help_button.grid(row=0, column=4, padx=(0, 22), pady=14)
        explain(help_button, 'Open the work instruction: what every button does, what the columns mean, and how to answer the questions this app is for.')

    def _build_update_banner(self):
        self.update_banner = ctk.CTkFrame(self, corner_radius=8, fg_color=('#fff4ce', '#4a3f1a'))
        self.update_banner.grid(row=1, column=0, sticky='ew', padx=22, pady=(0, 4))
        self.update_banner.grid_columnconfigure(0, weight=1)
        self.update_banner.grid_remove()
        self.update_label = ctk.CTkLabel(self.update_banner, text='', anchor='w', text_color=('#6b5300', '#ffe08a'), font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.update_label.grid(row=0, column=0, sticky='w', padx=14, pady=8)
        ctk.CTkButton(self.update_banner, text='Restart to update', width=140, height=28, command=self.restart_for_update, font=ctk.CTkFont(family=FONT, size=12)).grid(row=0, column=1, padx=(0, 6), pady=8)
        ctk.CTkButton(self.update_banner, text='Later', width=70, height=28, fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), command=lambda : self.update_banner.grid_remove(), font=ctk.CTkFont(family=FONT, size=12)).grid(row=0, column=2, padx=(0, 14), pady=8)

    def _build_stale_banner(self):
        self.stale_banner = ctk.CTkFrame(self, corner_radius=8, fg_color=('#fdeaea', '#4a1f22'))
        self.stale_banner.grid(row=2, column=0, sticky='ew', padx=22, pady=(0, 4))
        self.stale_banner.grid_columnconfigure(0, weight=1)
        self.stale_banner.grid_remove()
        self.stale_label = ctk.CTkLabel(self.stale_banner, text='', anchor='w', justify='left', text_color=('#8a1c24', '#ffb3bb'), font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.stale_label.grid(row=0, column=0, sticky='ew', padx=14, pady=8)
        self.stale_banner.bind('<Configure>', lambda event: self.stale_label.configure(wraplength=max(200, event.width - 140)))
        ctk.CTkButton(self.stale_banner, text='Dismiss', width=80, height=28, fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), command=lambda : self.stale_banner.grid_remove(), font=ctk.CTkFont(family=FONT, size=12)).grid(row=0, column=1, padx=(0, 14), pady=8)

    def _check_index_age(self):
        days = stale_index_age(self.index)
        if days is None:
            return
        log_activity(f'Index is {days} days old')
        self.post(lambda : self._show_stale_banner(days))

    def _show_stale_banner(self, days):
        built = self.index.built_on()
        when = f'{days} days ago, on {built}' if built else f'{days} days ago'
        self.stale_label.configure(text=f'This index was built {when}. Anything filed since then is missing from every answer here - searches and revision checklists will look normal and come up short. It should rebuild daily, so something has stopped: let the maintainer know.')
        self.stale_banner.grid()

    def _check_for_update(self):
        try:
            published = demo_config.published_exe()
            running = demo_config.running_exe()
            if not published or not running:
                return
            mine = os.stat(running)
            theirs = os.stat(published)
            if int(theirs.st_mtime) <= int(mine.st_mtime) and theirs.st_size == mine.st_size:
                return
            build = (int(theirs.st_mtime), theirs.st_size)
            if build == self._update_offered:
                return
            self._update_offered = build
            log_activity('Update available on the demo folder')
            self.post(self._show_update_banner)
        except OSError:
            pass

    def _schedule_update_check(self):
        delay = int(demo_config.UPDATE_CHECK_HOURS * 3600 * 1000)
        self._update_timer = self.after(delay, lambda : threading.Thread(target=self._poll_for_update, daemon=True).start())

    def _poll_for_update(self):
        self._check_for_update()
        self.post(self._schedule_update_check)

    def _show_update_banner(self):
        self.update_label.configure(text='  A newer version of BOM Search Demo is ready. It will install the next time you start the app.')
        self.update_banner.grid()

    def restart_for_update(self):
        launcher = os.environ.get('DEMO_LAUNCHER')
        if not launcher:
            share = os.environ.get('DEMO_SHARE') or demo_config.SHARE_APP
            launcher = os.path.join(os.path.dirname(share), 'BOM Search Demo.bat')
        log_activity('Restarting to take an update')
        try:
            if os.path.isfile(launcher):
                os.startfile(launcher)
            else:
                messagebox.showinfo('Restart to update', 'Close BOM Search Demo and start it again from this folder - the new version installs on the way in.')
                return
        except OSError as error:
            messagebox.showerror('Could not restart', str(error))
            return
        self._on_close()

    def _build_search(self):
        bar = ctk.CTkFrame(self, corner_radius=0, fg_color='transparent')
        bar.grid(row=3, column=0, sticky='ew', padx=22, pady=(14, 10))
        bar.grid_columnconfigure(0, weight=1)
        self.search_var = ctk.StringVar()
        self.search_box = ctk.CTkComboBox(bar, variable=self.search_var, height=46, font=ctk.CTkFont(family=FONT, size=17), dropdown_font=ctk.CTkFont(family=FONT, size=13), values=self.settings.get('recent', []), command=lambda choice: self.start_search())
        self.search_box.grid(row=0, column=0, sticky='ew', padx=(0, 10))
        self.search_box.bind('<Return>', lambda event: self.start_search())
        self.mode = self.settings.get('mode', 'boms')
        if self.mode not in MODE_COLUMNS:
            self.mode = 'boms'
        self.feedback_panel = None
        toggle = ctk.CTkFrame(bar, fg_color='transparent')
        toggle.grid(row=0, column=1, padx=(0, 10))
        self.mode_label_boms = ctk.CTkLabel(toggle, text=MODE_LABELS['boms'], font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.mode_label_boms.grid(row=0, column=0, padx=(0, 8))
        self.mode_switch = ctk.CTkSwitch(toggle, text='', width=52, command=self._mode_changed, onvalue='components', offvalue='boms')
        self.mode_switch.grid(row=0, column=1)
        explain(self.mode_switch, 'BOMs finds whole assemblies - one row per workbook.\nComponents finds individual parts - one row per part number.\n\nSame search box either way; a different kind of answer.')
        self.mode_label_components = ctk.CTkLabel(toggle, text=MODE_LABELS['components'], font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.mode_label_components.grid(row=0, column=2, padx=(8, 0))
        if self.mode == 'components':
            self.mode_switch.select()
        else:
            self.mode_switch.deselect()
        self.search_button = ctk.CTkButton(bar, text='Search', width=132, height=46, command=self.start_search, font=ctk.CTkFont(family=FONT, size=15, weight='bold'))
        self.search_button.grid(row=0, column=3)
        self.hint = ctk.CTkLabel(bar, text='', font=ctk.CTkFont(family=FONT, size=11), text_color=('gray50', 'gray55'))
        self.hint.grid(row=1, column=0, sticky='w', pady=(8, 0))
        quick = ctk.CTkFrame(bar, fg_color='transparent')
        quick.grid(row=1, column=1, columnspan=3, sticky='e', pady=(6, 0))
        ctk.CTkLabel(quick, text='Go straight to a part:', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray45', 'gray60')).grid(row=0, column=0, padx=(0, 8))
        self.quick_var = ctk.StringVar()
        self.quick_box = ctk.CTkEntry(quick, textvariable=self.quick_var, width=140, height=30, placeholder_text='part number', font=ctk.CTkFont(family=FONT, size=12))
        self.quick_box.grid(row=0, column=1)
        self.quick_box.bind('<Return>', lambda event: self.quick_open('drawing'))
        self.bind_all('<KeyPress>', lambda event: setattr(self, '_last_input', 'key'), add='+')
        self.bind_all('<ButtonPress>', lambda event: setattr(self, '_last_input', 'mouse'), add='+')
        self.quick_drawing = ctk.CTkButton(quick, text='Its drawing', width=106, height=30, command=lambda : self.quick_open('drawing'), fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=12))
        self.quick_drawing.grid(row=0, column=2, padx=(6, 0))
        self.quick_cad = ctk.CTkButton(quick, text='Its CAD', width=90, height=30, command=lambda : self.quick_open('cad'), fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=12))
        self.quick_cad.grid(row=0, column=3, padx=(6, 0))
        self._update_hint()

    def _update_hint(self):
        common = 'Ctrl+F to jump here   ·   Esc to clear   ·   click a heading to sort   ·   Alt+click a cell to search it   ·   right-click for more'
        if self.mode == 'components':
            self.hint.configure(text='Showing parts, not assemblies   ·   double-click to open the part’s drawing   ·   ' + common)
        else:
            self.hint.configure(text='double-click to open a BOM   ·   ' + common)

    def current_scope(self):
        return 'all'

    def _mode_changed(self):
        self.mode = 'components' if self.mode_switch.get() == 'components' else 'boms'
        if getattr(self, 'columns_panel', None) is not None:
            self.columns_panel.close()
        self.settings['mode'] = self.mode
        save_settings(self.settings)
        self._configure_table()
        self._apply_accent()
        self._update_hint()
        self.hide_components()
        self.results = None
        self._term = ''
        self.table.delete(*self.table.get_children())
        self._set_actions('disabled')
        self.export_button.configure(state='disabled')
        self.commodity_tabs.configure(values=['All'])
        self.commodity_tabs.set('All')
        what = MODE_LABELS[self.mode].lower()
        self.results_label.configure(text=f'Searching {what} — press Search')
        if self.search_var.get().strip():
            self.set_status(f'Now searching {what}. Press Search to run “{self.search_var.get().strip()}” again.')
        else:
            self.set_status(f'Now searching {what}.')

    def _build_suggestion_bar(self, wrapper):
        self.suggestion_bar = ctk.CTkFrame(wrapper, corner_radius=8, fg_color=('#fdf1d6', '#3b3320'))
        self.suggestion_bar.grid(row=2, column=0, sticky='ew', padx=12, pady=(0, 8))
        self.suggestion_bar.grid_columnconfigure(5, weight=1)
        self.suggestion_label = ctk.CTkLabel(self.suggestion_bar, text='', anchor='w', justify='left', font=ctk.CTkFont(family=FONT, size=12))
        self.suggestion_label.grid(row=0, column=0, sticky='w', padx=14, pady=(9, 9))
        self.suggestion_buttons = []
        self.suggestion_bar.grid_remove()

    def _suggest(self, text, buttons):
        for widget in self.suggestion_buttons:
            widget.destroy()
        self.suggestion_buttons = []
        if not buttons:
            self.suggestion_bar.grid_remove()
            return
        self.suggestion_label.configure(text=text)
        for (position, (caption, command)) in enumerate(buttons[:4]):
            button = ctk.CTkButton(self.suggestion_bar, text=caption, height=28, width=max(96, 9 * len(caption)), font=ctk.CTkFont(family=FONT, size=12), command=command)
            button.grid(row=0, column=position + 1, padx=(6, 0), pady=8)
            self.suggestion_buttons.append(button)
        self.suggestion_bar.grid()

    def _research(self, term):
        self.search_var.set(term)
        self.start_search()

    def _show_suggestions(self, term, results):
        if self.index is None:
            return
        found = results is not None and (not results.empty)
        drawings = self.index.drawings_matching(term)
        buttons = []
        if found:
            if not drawings.empty and len(drawings) <= 2:
                for (_, row) in drawings.iterrows():
                    name = str(row['PDF_Name'])
                    buttons.append((os.path.splitext(name)[0], lambda folder=str(row['Folder']), pdf=name: self._open_drawing_file(folder, pdf)))
                text = self._drawing_question(term)
            if self.mode == 'boms' and len(term.strip()) == 10 and term.strip().isdigit():
                parts = self.index.components()
                if not parts.empty and (parts['part_number'].astype(str) == term.strip()).any():
                    buttons.append(('Show the part itself', lambda : self._switch_to_components(term)))
                    text = self._drawing_question(term) if len(buttons) > 1 else f'“{term}” is a part number. These are the BOMs that use it.'
            if not buttons:
                self.suggestion_bar.grid_remove()
                return
            self._suggest(text, buttons)
            return
        if not drawings.empty:
            for (_, row) in drawings.iterrows():
                name = str(row['PDF_Name'])
                buttons.append((os.path.splitext(name)[0], lambda folder=str(row['Folder']), pdf=name: self._open_drawing_file(folder, pdf)))
            self._suggest(f'No BOM lists “{term}”. Did you want to open its drawing?', buttons)
            log_activity(f"Offered a drawing for '{term}' after no BOM matched")
            return
        alternatives = self.index.alternative_terms(term)
        if alternatives:
            for (word, count) in alternatives:
                buttons.append((f'{word}  ({count})', lambda t=word: self._research(t)))
            self._suggest(f'Nothing matched “{term}” exactly. Try:', buttons)
            log_activity(f"Suggested other words for '{term}'")
            return
        near = self.index.similar_numbers(term)
        if near:
            for number in near[:4]:
                buttons.append((number, lambda t=number: self._research(t)))
            self._suggest(f'Nothing matched “{term}”. Did you mean:', buttons)
            log_activity(f"Suggested near numbers for '{term}'")
            return
        self.suggestion_bar.grid_remove()

    def _drawing_question(self, term):
        wanted = term.strip()[:demo_index.DRAWING_NUMBER_LENGTH]
        numbers = self.index.files.get('BOM #')
        if wanted.isdigit() and numbers is not None and numbers.astype(str).str.startswith(wanted).any():
            return 'It looks like you are searching for an assembly. Did you want to open its drawing?'
        return 'It looks like you are searching for a component. Did you want to open its drawing?'

    def _switch_to_components(self, term):
        self.mode = 'components'
        self.mode_switch.select()
        self.settings['mode'] = self.mode
        save_settings(self.settings)
        self._configure_table()
        self._apply_accent()
        self._update_hint()
        self.hide_components()
        self._research(term)

    def _open_drawing_file(self, folder, pdf_name):
        path = os.path.join(demo_config.paths()['drawings_root'], demo_config.drawings_folder_for(folder), pdf_name)
        self._open(path, f'drawing {pdf_name} (no BOM lists it)')

    def choose_columns(self):
        if getattr(self, 'columns_panel', None) is not None:
            self.columns_panel.close()
            return
        self.columns_panel = ColumnsPanel(self)

    def visible_columns(self, mode=None):
        mode = mode or self.mode
        keys = [key for (key, *_) in MODE_COLUMNS[mode]]
        hidden = (self.settings.get('hidden_columns') or {}).get(mode) or []
        kept = [key for key in keys if key not in hidden]
        return kept or keys

    def show_columns(self, keys):
        hidden = dict(self.settings.get('hidden_columns') or {})
        hidden[self.mode] = [key for (key, *_) in MODE_COLUMNS[self.mode] if key not in keys]
        self.settings['hidden_columns'] = hidden
        save_settings(self.settings)
        self._configure_table()
        self._render_rows()
        hidden = len(MODE_COLUMNS[self.mode]) - len(self.visible_columns())
        self.set_status('Showing every column.' if not hidden else f"{hidden} column{('s' if hidden != 1 else '')} hidden. The export follows what is on screen.")
        log_activity(f'Columns in {self.mode}: {len(self.visible_columns())} shown')

    def _configure_table(self):
        columns = MODE_COLUMNS[self.mode]
        self.table.configure(columns=[key for (key, *_) in columns], displaycolumns=self.visible_columns(), show='headings')
        for (key, heading, width, anchor) in columns:
            self.table.heading(key, text=heading, command=lambda k=key: self.sort_by(k))
            self.table.column(key, width=width, minwidth=min(width, 80), anchor=anchor)
        self._sort_state = {}

    def _build_results(self):
        wrapper = self.results_panel = ctk.CTkFrame(self, corner_radius=10)
        wrapper.grid(row=4, column=0, sticky='nsew', padx=22, pady=(0, 10))
        wrapper.grid_rowconfigure(3, weight=1)
        wrapper.grid_columnconfigure(0, weight=1)
        toolbar = ctk.CTkFrame(wrapper, fg_color='transparent')
        toolbar.grid(row=0, column=0, sticky='ew', padx=12, pady=(10, 6))
        toolbar.grid_columnconfigure(0, weight=1)
        self.results_label = ctk.CTkLabel(toolbar, text='No search yet', font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.results_label.grid(row=0, column=0, sticky='w')
        self.commodity_tabs = ctk.CTkSegmentedButton(wrapper, values=['All'], command=lambda choice: self._render_rows(), font=ctk.CTkFont(family=FONT, size=11), height=32, border_width=3, corner_radius=6, fg_color=('gray72', 'gray28'), unselected_color=('gray90', 'gray20'), unselected_hover_color=('gray82', 'gray32'), selected_color=('#1a6fd4', '#2f6fb5'), selected_hover_color=('#1a6fd4', '#3b7fc4'), text_color=('gray10', 'gray90'))
        self.commodity_tabs.set('All')
        self.commodity_tabs.grid(row=1, column=0, sticky='w', padx=12, pady=(0, 8))
        explain(self.commodity_tabs, 'Narrow the results to one commodity. The number on each tab is how many of your results are in it.')
        self._tab_labels = {'All': 'All'}
        self.action_buttons = []
        self.action_labels = []
        for (index, (label, command)) in enumerate((('Open BOM', self.open_bom), ('Drawing', self.open_drawing), ('Customer Dwg', self.open_customer_drawing), ('DFMEA', self.open_dfmea), ('CAD', self.open_cad), ('Components', self.show_components), ('Projects', self.show_projects))):
            self.action_labels.append(label)
            width = 104
            if label in ROW_WHAT:
                command = lambda opener=command, what=ROW_WHAT[label]: self._for_rows(opener, what=what)
            button = ctk.CTkButton(toolbar, text=label, width=width, height=30, command=command, state='disabled', border_width=1, fg_color=ROW_BUTTON_OFF, border_color=ROW_BUTTON_OFF, text_color_disabled=('gray45', 'gray60'), font=ctk.CTkFont(family=FONT, size=12))
            button.grid(row=0, column=index + 2, padx=(6, 0))
            explain(button, ACTION_HELP.get(label, ''))
            self.action_buttons.append(button)
        self.export_button = ctk.CTkButton(toolbar, text='Export', width=84, height=30, command=self.export_results, state='disabled', fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=12))
        self.export_button.grid(row=0, column=len(self.action_buttons) + 2, padx=(10, 0))
        explain(self.export_button, 'Save everything currently listed to an Excel file - the whole result set, not just the highlighted row.')
        self.columns_button = ctk.CTkButton(toolbar, text='Columns', width=90, height=30, command=self.choose_columns, fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=12))
        self.columns_button.grid(row=0, column=len(self.action_buttons) + 3, padx=(6, 0))
        explain(self.columns_button, 'Choose which columns this view shows. Kept for next time, separately for BOMs and for parts, and the export follows it.')
        ctk.CTkLabel(toolbar, text='Selected row:', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray45', 'gray60')).grid(row=0, column=1, padx=(0, 8))
        self._build_suggestion_bar(wrapper)
        holder = tk.Frame(wrapper, highlightthickness=0, bd=0)
        holder.grid(row=3, column=0, sticky='nsew', padx=12, pady=(0, 12))
        holder.grid_rowconfigure(0, weight=1)
        holder.grid_columnconfigure(0, weight=1)
        self.table_holder = holder
        self.table = ttk.Treeview(holder, columns=[key for (key, *_) in MODE_COLUMNS[self.mode]], show='headings', selectmode='extended', style='Demo.Treeview')
        self.table.grid(row=0, column=0, sticky='nsew')
        self._sort_state = {}
        self._configure_table()
        scroll = ttk.Scrollbar(holder, orient='vertical', command=self.table.yview)
        sideways = ttk.Scrollbar(holder, orient='horizontal', command=self.table.xview)
        self.table.configure(yscrollcommand=scroll.set, xscrollcommand=sideways.set)
        scroll.grid(row=0, column=1, sticky='ns')
        sideways.grid(row=1, column=0, sticky='ew')
        HeadingTooltip(self.table, COLUMN_HELP)
        self.table.bind('<<TreeviewSelect>>', self._on_row_selected)
        self.table.bind('<Double-1>', lambda event: self.open_drawing() if self.mode == 'components' else self.open_bom())
        self.table.bind('<Alt-Button-1>', self._search_clicked_cell)
        self.table.bind('<Button-3>', self._show_context_menu)
        for sequence in ('<Control-a>', '<Control-A>'):
            self.table.bind(sequence, lambda event: self._select_all(self.table))
        self._sort_state = {}
        self.context_menu = tk.Menu(self, tearoff=0)
        for (label, command) in (('Open BOM', lambda : self._for_rows(self.open_bom, what='BOM')), ('Open drawing', lambda : self._for_rows(self.open_drawing, what='drawing')), ('Open customer drawing', lambda : self._for_rows(self.open_customer_drawing, what='customer drawing')), ('Open CAD', lambda : self._for_rows(self.open_cad, what='CAD file'))):
            self.context_menu.add_command(label=label, command=command)
        self.context_menu.add_separator()
        for (label, command) in (('Copy BOM file', lambda : self._for_rows(self.open_bom, act=self._copy_file, what='BOM')), ('Copy drawing file', lambda : self._for_rows(self.open_drawing, act=self._copy_file, what='drawing')), ('Copy CAD file', lambda : self._for_rows(self.open_cad, act=self._copy_file, what='CAD file')), ('Copy DFMEA file', lambda : self._for_rows(self.open_dfmea, act=self._copy_file, what='DFMEA'))):
            self.context_menu.add_command(label=label, command=command)
        self.context_menu.add_separator()
        self.context_menu.add_command(label='Components in this BOM', command=self.show_components)
        self.context_menu.add_separator()
        self.context_menu.add_command(label='Copy BOM number', command=lambda : self._copy(False))
        self.context_menu.add_command(label='Copy whole row', command=lambda : self._copy(True))
        self.component_menu = tk.Menu(self, tearoff=0)
        for (label, command) in (('Open drawing', lambda : self._for_rows(self.open_drawing, what='drawing')), ('Open CAD', lambda : self._for_rows(self.open_cad, what='CAD file'))):
            self.component_menu.add_command(label=label, command=command)
        self.component_menu.add_separator()
        for (label, command) in (('Copy drawing file', lambda : self._for_rows(self.open_drawing, act=self._copy_file, what='drawing')), ('Copy CAD file', lambda : self._for_rows(self.open_cad, act=self._copy_file, what='CAD file'))):
            self.component_menu.add_command(label=label, command=command)
        self.component_menu.add_separator()
        self.component_menu.add_command(label='Open a BOM that uses it', command=lambda : self._for_rows(self.open_bom, what='BOM'))
        self.component_menu.add_command(label='Find every BOM that uses it', command=self.find_boms_using_part)
        self.component_menu.add_command(label='What else would need revising?', command=self.show_revision_impact)
        self.component_menu.add_separator()
        self.component_menu.add_command(label='Copy part number', command=lambda : self._copy(False))
        self.component_menu.add_command(label='Copy whole row', command=lambda : self._copy(True))
        self.parts_menu = tk.Menu(self, tearoff=0)
        for (label, command) in (('Open drawing', lambda : self._for_parts(self.open_part_drawing, what='drawing')), ('Open CAD', lambda : self._for_parts(self.open_part_cad, what='CAD file'))):
            self.parts_menu.add_command(label=label, command=command)
        self.parts_menu.add_separator()
        for (label, command) in (('Copy drawing file', lambda : self._for_parts(self.open_part_drawing, act=self._copy_file, what='drawing')), ('Copy CAD file', lambda : self._for_parts(self.open_part_cad, act=self._copy_file, what='CAD file'))):
            self.parts_menu.add_command(label=label, command=command)
        self.parts_menu.add_separator()
        self.parts_menu.add_command(label='Where else is it used?', command=self.search_selected_part)
        self.parts_menu.add_command(label='What if it revs?', command=self.show_revision_impact)
        self.parts_menu.add_separator()
        self.parts_menu.add_command(label='Copy part number', command=self.copy_part_number)

    def _build_components(self):
        self.components = ctk.CTkFrame(self, corner_radius=10)
        self.components.grid(row=5, column=0, sticky='nsew', padx=22, pady=(0, 10))
        self.components.grid_columnconfigure(0, weight=1)
        self.components.grid_rowconfigure(1, weight=1)
        self.components.grid_remove()
        bar = ctk.CTkFrame(self.components, fg_color='transparent')
        bar.grid(row=0, column=0, sticky='ew', padx=12, pady=(10, 6))
        bar.grid_columnconfigure(2, weight=1)
        self.components_label = ctk.CTkLabel(bar, text='', font=ctk.CTkFont(family=FONT, size=13, weight='bold'))
        self.components_label.grid(row=0, column=0, sticky='w', padx=(0, 16))
        filter_group = ctk.CTkFrame(bar, fg_color='transparent')
        filter_group.grid(row=0, column=1, sticky='w')
        ctk.CTkLabel(filter_group, text='Filter this list:', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray35', 'gray70')).pack(side='left', padx=(0, 8))
        self.parts_filter_var = ctk.StringVar()
        self.parts_filter = ctk.CTkEntry(filter_group, textvariable=self.parts_filter_var, width=230, height=28, placeholder_text='part number, name, material...', font=ctk.CTkFont(family=FONT, size=12))
        self.parts_filter.pack(side='left')
        self.parts_filter_var.trace_add('write', lambda *a: self._render_parts())
        self.parts_filter.bind('<Escape>', lambda event: self.parts_filter_var.set(''))
        for (column, (label, width, command)) in enumerate((('Open drawing', 110, lambda : self._for_parts(self.open_part_drawing, what='drawing')), ('Open CAD', 90, lambda : self._for_parts(self.open_part_cad, what='CAD file')), ('Where else is it used?', 170, self.search_selected_part), ('What if it revs?', 130, self.show_revision_impact), ('Close', 70, self.hide_components)), start=3):
            button = ctk.CTkButton(bar, text=label, width=width, height=28, command=command, font=ctk.CTkFont(family=FONT, size=12))
            button.grid(row=0, column=column, padx=(0, 6))
            if label == 'Where else is it used?':
                self.parts_where_button = button
                explain(button, 'Search for every BOM that uses this part. One row only.')
            if label == 'What if it revs?':
                self.parts_revision_button = button
                explain(button, 'What needs revising if this part changes - the same checklist as the row button. Off on a prototype row: this is for released BOMs. One row only.')
        holder = tk.Frame(self.components, highlightthickness=0, bd=0)
        holder.grid(row=1, column=0, sticky='nsew', padx=12, pady=(0, 12))
        holder.grid_rowconfigure(0, weight=1)
        holder.grid_columnconfigure(0, weight=1)
        self.parts_table = ttk.Treeview(holder, columns=[key for (key, _) in demo_index.PART_COLUMNS], show='headings', selectmode='extended', style='Demo.Treeview', height=8)
        for (key, heading) in demo_index.PART_COLUMNS:
            self.parts_table.heading(key, text=heading, command=lambda k=key: self.sort_parts_by(k))
            width = 90 if key in ('level', 'quantity') else 280 if key == 'part_name' else 140
            self.parts_table.column(key, width=width, minwidth=min(width, 80), anchor='center')
        self.parts_table.grid(row=0, column=0, sticky='nsew')
        self.parts_table.bind('<Double-1>', lambda event: self.open_part_drawing())
        self.parts_table.bind('<Button-3>', self._show_parts_menu)
        self.parts_table.bind('<<TreeviewSelect>>', self._on_part_selected)
        for sequence in ('<Control-a>', '<Control-A>'):
            self.parts_table.bind(sequence, lambda event: self._select_all(self.parts_table))
        scroll = ttk.Scrollbar(holder, orient='vertical', command=self.parts_table.yview)
        sideways = ttk.Scrollbar(holder, orient='horizontal', command=self.parts_table.xview)
        self.parts_table.configure(yscrollcommand=scroll.set, xscrollcommand=sideways.set)
        scroll.grid(row=0, column=1, sticky='ns')
        sideways.grid(row=1, column=0, sticky='ew')

    def _build_status(self):
        self.status = ctk.CTkLabel(self, text='Starting...', anchor='w', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray40', 'gray65'))
        self.status.grid(row=6, column=0, sticky='ew', padx=24, pady=(0, 12))

    def _style_table(self):
        mode = 'dark' if ctk.get_appearance_mode() == 'Dark' else 'light'
        colours = dict(TABLE_COLOURS[mode])
        colours['selected'] = MODE_ACCENTS[self.mode]['table'][mode]
        style = ttk.Style()
        style.theme_use('clam')
        style.configure('Demo.Treeview', background=colours['bg'], fieldbackground=colours['bg'], foreground=colours['fg'], borderwidth=0, rowheight=30, font=(FONT, 11))
        style.configure('Demo.Treeview.Heading', background=colours['heading_bg'], foreground=colours['heading_fg'], relief='flat', borderwidth=0, padding=(8, 8), font=(FONT, 11, 'bold'))
        style.map('Demo.Treeview.Heading', background=[('active', colours['selected'])])
        style.map('Demo.Treeview', background=[('selected', colours['selected'])], foreground=[('selected', colours['selected_fg'])])
        for bar in ('Vertical.TScrollbar', 'Horizontal.TScrollbar'):
            style.configure(bar, background=colours['heading_bg'], troughcolor=colours['bg'], bordercolor=colours['bg'], arrowcolor=colours['heading_fg'])
        for table in (getattr(self, 'table', None), getattr(self, 'parts_table', None)):
            if table is not None:
                table.tag_configure('prototype', background=colours['prototype'])
                table.tag_configure('odd', background=colours['alt'])
                table.tag_configure('even', background=colours['bg'])
        for holder in (getattr(self, 'table_holder', None),):
            if holder is not None:
                holder.configure(background=colours['bg'])

    def _apply_accent(self):
        accent = MODE_ACCENTS[self.mode]
        (normal, hover) = (accent['normal'], accent['hover'])
        self.search_button.configure(fg_color=normal, hover_color=hover)
        for button in self.action_buttons:
            button.configure(**self._action_look(button.cget('state')))
        self.mode_switch.configure(progress_color=normal)
        self.commodity_tabs.configure(selected_color=normal, selected_hover_color=hover)
        self.results_panel.configure(fg_color=accent['panel'])
        live = ('gray10', 'gray95')
        dim = ('gray60', 'gray50')
        self.mode_label_boms.configure(text_color=live if self.mode == 'boms' else dim)
        self.mode_label_components.configure(text_color=live if self.mode == 'components' else dim)
        self._style_table()

    def _on_theme_change(self, value):
        ctk.set_appearance_mode(value.lower())
        self.settings['theme'] = value.lower()
        save_settings(self.settings)
        self._style_table()

    def post(self, func):
        self.ui_queue.put(func)

    def _drain_ui_queue(self):
        try:
            while True:
                self.ui_queue.get_nowait()()
        except queue.Empty:
            pass
        except Exception:
            pass
        self.after(50, self._drain_ui_queue)

    def set_status(self, text):
        self.post(lambda : self.status.configure(text=text))

    def _startup(self):
        paths = demo_config.paths()
        state = network_state(paths)
        if state != 'ready':
            self.index_error = state
            self.set_status(network_status_text(state))
            log_activity(f'Could not start: {state}')
            self.post(lambda : self._show_network_card(state))
            return
        try:
            wanted = {'projects': paths.get('projects'), 'project folders': paths.get('project_folders'), 'obsolete BOMs': demo_config.obsolete_index()}
            (chosen, sampled) = ({}, set())
            for (name, path) in wanted.items():
                (found, is_sample) = demo_config.resolve_table(path)
                chosen[name] = found
                if is_sample:
                    sampled.add(name)
            self.index = demo_index.load(paths['cells'], paths['files'], paths['pdfs'], projects_path=chosen['projects'], project_folders_path=chosen['project folders'], obsolete_path=chosen['obsolete BOMs'], dfmea_path=paths.get('dfmea'), prototype_dir=paths.get('prototype_index'))
            self.index.sampled = sampled
        except demo_index.IndexError_ as error:
            self.index_error = error
            self.set_status('The index could not be loaded.')
            self.post(lambda : messagebox.showerror('Index unavailable', str(error)))
            return
        except Exception as error:
            self.index_error = error
            self.set_status('The index could not be loaded.')
            self.post(lambda : messagebox.showerror('Index unavailable', str(error)))
            return
        self.set_status(f'Ready. {self.index.describe()}{prototype_note(self.index)}')
        self._check_index_age()
        self._check_for_update()
        self.post(self._schedule_update_check)
        try:
            self.index.components()
        except Exception as error:
            log_activity(f'Component table could not be prepared: {error}')
        try:
            self.index.warm_revision_walk()
        except Exception as error:
            log_activity(f'Assembly trees could not be prepared: {error}')
        if self.index.prototype is not None:
            try:
                self.index.prototype.components()
                self.index.prototype.warm_revision_walk()
            except Exception as error:
                log_activity(f'Prototype index could not be prepared: {error}')
                self.index.prototype_error = f'{type(error).__name__}: {error}'
                self.index.prototype = None

    def _show_network_card(self, state):
        if self.network_card is not None:
            self.network_card.close()
        self.network_card = NetworkCard(self, state, self.load_index_again)

    def load_index_again(self):
        self.set_status('Looking for the index again...')
        threading.Thread(target=self._startup, daemon=True).start()

    def _focus_search(self):
        self.search_box.focus_set()

    def start_search(self):
        term = self.search_var.get().strip()
        if not term:
            self.set_status('Type something to search for.')
            return
        if self.index is None:
            if self.index_error in ('offline', 'no index'):
                self.set_status(network_status_text(self.index_error))
            else:
                self.set_status('Still loading the index...' if not self.index_error else 'The index could not be loaded.')
            return
        if getattr(self, 'projects_panel', None) is not None:
            self.projects_panel.close()
        scope = self.current_scope()
        mode = self.mode
        self._remember(term)
        log_activity(f'Searched for: {term} ({MODE_LABELS[mode]})' + ('' if scope == 'all' else f' (scope: {scope})'))
        self.search_button.configure(state='disabled', text='Searching')
        self.set_status(f'Searching {MODE_LABELS[mode].lower()} for {term}...')
        threading.Thread(target=self._run_search, args=(term, scope, mode), daemon=True).start()

    def _run_search(self, term, scope, mode):
        started = datetime.now()
        try:
            if mode == 'components':
                results = demo_index.search_components_both(self.index, term, scope=scope)
            else:
                results = demo_index.search_both(self.index, term, scope=scope)
        except Exception as error:
            self.set_status(f'Search failed: {error}')
            self.post(lambda : self.search_button.configure(state='normal', text='Search'))
            return
        elapsed = (datetime.now() - started).total_seconds() * 1000
        self.post(lambda : self._show_results(term, results, elapsed, scope))

    def _show_results(self, term, results, elapsed, scope='all', caption=None):
        self.results = results
        self._results_caption = caption
        self._sort_state = {}
        self._mark_sorted()
        self.table.delete(*self.table.get_children())
        self.hide_components()
        if results is None or results.empty:
            if scope != 'all':
                where = demo_index.SCOPE_LABELS.get(scope, scope)
                self.results_label.configure(text=f'No results for “{term}” in {where}')
                self.set_status(f'Nothing matched “{term}” in {where} ({elapsed:,.0f} ms). Set the dropdown to Everything to search all columns.{prototype_note(self.index)}')
            else:
                self.results_label.configure(text=f'No results for “{term}”')
                self.set_status(f'Nothing matched “{term}” ({elapsed:,.0f} ms).{prototype_note(self.index)}')
            self._show_suggestions(term, results)
            log_activity(f"Result: '{term}' -> NO RESULTS in {elapsed:.0f} ms" + ('' if scope == 'all' else f' (scope: {scope})'))
            self._set_actions('disabled')
            self.export_button.configure(state='disabled')
            self.commodity_tabs.configure(values=['All'])
            self.commodity_tabs.set('All')
            self.search_button.configure(state='normal', text='Search')
            return
        self._show_suggestions(term, results)
        counts = commodity_counts(results)
        self._tab_labels = {'All': 'All'}
        labels = [f'All  ({len(results):,})']
        self._tab_labels[labels[0]] = 'All'
        for commodity in sorted(counts.index):
            short = demo_config.commodity_name(commodity)
            label = f'{short}  ({counts[commodity]})'
            labels.append(label)
            self._tab_labels[label] = commodity
        if 'Source' in results.columns:
            prototypes = int((results['Source'] == demo_index.PROTOTYPE).sum())
            if prototypes:
                label = f'Prototype  ({prototypes})'
                labels.append(label)
                self._tab_labels[label] = PROTOTYPE_TAB
        self.commodity_tabs.configure(values=labels)
        self.commodity_tabs.set(labels[0])
        self._term = term
        self._render_rows()
        where = '' if scope == 'all' else f' in {demo_index.SCOPE_LABELS.get(scope, scope)}'
        what = 'parts' if self.mode == 'components' else 'BOMs'
        padded = getattr(self.index, 'zero_padded_hits', 0)
        note = ''
        if padded:
            note = f" {padded} cell{('s' if padded != 1 else '')} write the number without its leading zero - Excel drops it when the cell is a number - and are included."
        prototypes = int((results['Source'] == demo_index.PROTOTYPE).sum()) if 'Source' in results.columns else 0
        mix = f' ({prototypes:,} prototype)' if prototypes else ''
        self.set_status(f'Found {len(results):,} {what}{mix}{where} in {elapsed:,.0f} ms.{note} {self.index.describe()}{prototype_note(self.index)}')
        log_activity(f"Result: '{term}' -> {len(results)} rows in {elapsed:.0f} ms")
        self.export_button.configure(state='normal')
        self.search_button.configure(state='normal', text='Search')

    def _index_for(self, row):
        prototype = getattr(self.index, 'prototype', None)
        if prototype is not None and _is_prototype(row):
            return prototype
        return self.index

    def _folder_for(self, row, key):
        name = 'proto_' + key if _is_prototype(row) else key
        return demo_config.paths().get(name, '')

    def _bom_row_values(self, row):
        index = self._index_for(row)
        return (row.get('BOM #', ''), row.get('Drawing', ''), _one_line(row.get('Assembly Name:', '')), _one_line(row.get('Customer & Platform:', '')), _one_line(row.get('Customer Assembly No:', '')), _one_line(row.get('Customer Service No:', '')), folder_label(row), index.bom_plant(row.get('Filename', '')), index.bom_date(row.get('Filename', ''), row.get('ECN#:', '')), index.bom_signer(row.get('Filename', '')), row.get('MatchedIn', ''), _one_line(row.get('MatchedValue', ''), 160))

    def _component_row_values(self, row):
        return (row.get('part_number', ''), _one_line(row.get('part_name', '')), row.get('drawing', ''), _one_line(row.get('material', ''), 120), _thickness(row.get('thickness_min', '')), _thickness(row.get('thickness_max', '')), _one_line(row.get('finish', ''), 90), folder_label(row), row.get('BOMs', ''), row.get('MatchedIn', ''))

    def _row_values(self, row):
        return self._component_row_values(row) if self.mode == 'components' else self._bom_row_values(row)

    def _render_rows(self):
        if self.results is None or self.results.empty:
            return
        chosen = self._tab_labels.get(self.commodity_tabs.get(), 'All')
        self.table.delete(*self.table.get_children())
        components = self.mode == 'components'
        shown = 0
        for (position, (_, row)) in enumerate(self.results.iterrows()):
            if not _row_is_shown(row, chosen):
                continue
            values = self._row_values(row)
            self.table.insert('', 'end', iid=str(position), values=values, tags=self._row_tags(shown, row))
            shown += 1
        term = getattr(self, '_term', '')
        what = 'parts' if components else 'BOMs'
        caption = getattr(self, '_results_caption', None)
        name = demo_config.commodity_name(chosen)
        if caption:
            self.results_label.configure(text=caption if chosen == 'All' else f'{caption}  ·  {name} ({shown:,})')
        elif chosen == 'All':
            self.results_label.configure(text=f'{len(self.results):,} {what} matching “{term}”')
        else:
            self.results_label.configure(text=f'{shown:,} of {len(self.results):,} {what} matching “{term}”  ·  {name}')
        self._set_actions('disabled')

    def _row_tags(self, position, row):
        return self._tags_for(position, _is_prototype(row))

    def _tags_for(self, position, is_prototype):
        return ('odd' if position % 2 else 'even',) + (('prototype',) if is_prototype else ())

    def sort_by(self, key):
        if self.results is None or self.results.empty:
            return
        descending = not self._sort_state.get(key, False)
        self._sort_state = {key: descending}
        rows = [(self.table.set(item, key), item) for item in self.table.get_children('')]
        self._mark_sorted(key, descending)
        _sort_rows(rows, descending)
        prototypes = self.results['Source'].to_numpy() == demo_index.PROTOTYPE if 'Source' in self.results.columns else None
        for (position, (_, item)) in enumerate(rows):
            self.table.move(item, '', position)
            is_prototype = prototypes is not None and prototypes[int(item)]
            self.table.item(item, tags=self._tags_for(position, is_prototype))

    def _mark_sorted(self, sorted_key=None, descending=False):
        for (key, heading, *_) in MODE_COLUMNS[self.mode]:
            arrow = ''
            if key == sorted_key:
                arrow = '  ▾' if descending else '  ▴'
            try:
                self.table.heading(key, text=heading + arrow)
            except tk.TclError:
                pass

    def _remember(self, term):
        recent = [item for item in self.settings.get('recent', []) if item.lower() != term.lower()]
        recent.insert(0, term)
        self.settings['recent'] = recent[:10]
        save_settings(self.settings)
        self.search_box.configure(values=self.settings['recent'])

    def _on_row_selected(self, event=None):
        count = len(self.table.selection())
        self._set_actions('normal' if count else 'disabled')
        if count > 1:
            self.set_status(f'{count} rows selected.')

    def _select_all(self, table):

        def every(parent=''):
            for item in table.get_children(parent):
                yield item
                yield from every(item)
        table.selection_set(list(every()))
        return 'break'

    def _keep_or_select(self, table, row_id):
        if row_id not in table.selection():
            table.selection_set(row_id)
        table.focus(row_id)

    def _for_rows(self, opener, act=None, what='file'):
        self._many(self.table, list(self.table.selection()), opener, act, what)

    def _show_menu_states(self):
        rows = self.selected_rows()
        several = len(rows) > 1
        prototype = bool(rows) and all((_is_prototype(row) for row in rows))
        self.context_menu.entryconfigure('Open customer drawing', state='disabled' if prototype else 'normal')
        self.context_menu.entryconfigure('Components in this BOM', state='disabled' if several else 'normal')
        for label in ('Find every BOM that uses it', 'What else would need revising?'):
            off = several or (prototype and label.startswith('What'))
            self.component_menu.entryconfigure(label, state='disabled' if off else 'normal')

    def _action_look(self, state):
        if state == 'disabled':
            return dict(fg_color=ROW_BUTTON_OFF, border_color=ROW_BUTTON_OFF)
        accent = MODE_ACCENTS[self.mode]
        return dict(fg_color=accent['normal'], hover_color=accent['hover'], border_color=ROW_BUTTON_BORDER)

    def _set_actions(self, state):
        rows = self.selected_rows() if state == 'normal' else []
        prototype = bool(rows) and all((_is_prototype(row) for row in rows))
        for (label, button) in zip(self.action_labels, self.action_buttons):
            wanted = state
            if len(rows) > 1 and label in ONE_ROW_ACTIONS:
                wanted = 'disabled'
            elif prototype and label in PRODUCTION_ONLY_ACTIONS:
                wanted = 'disabled'
            elif state == 'normal' and self.mode == 'components' and (label in ('Components', 'DFMEA')):
                wanted = 'disabled'
            button.configure(state=wanted, **self._action_look(wanted))

    def selected_row(self):
        batch = getattr(self, '_batch', None)
        if batch is not None and batch.owner is self.table:
            return self.results.iloc[int(batch.item)]
        selection = self.table.selection()
        if not selection or self.results is None:
            return None
        return self.results.iloc[int(selection[0])]

    def selected_rows(self):
        if self.results is None:
            return []
        return [self.results.iloc[int(item)] for item in self.table.selection()]

    def _missed(self, number):
        batch = getattr(self, '_batch', None)
        if batch is not None:
            batch.missing.append(str(number).strip() or '(no number)')

    def _say(self, title, text, number='', error=False, note=None):
        batch = getattr(self, '_batch', None)
        if batch is None:
            (messagebox.showerror if error else messagebox.showinfo)(title, text)
        elif note:
            batch.note(note)
        else:
            self._missed(number)

    def _many(self, owner, items, opener, act=None, what='file'):
        if len(items) <= 1:
            opener(act=act) if act is not None else opener()
            return
        batch = Batch(owner, copying=act is not None)
        self._batch = batch
        try:
            for item in items:
                batch.item = item
                try:
                    opener(act=batch.collect)
                except Exception as error:
                    batch.note(f'One row could not be read: {error}')
        finally:
            self._batch = None
        self._finish_batch(batch, what, len(items))

    def _finish_batch(self, batch, what, rows):
        files = [path for path in batch.files if os.path.exists(path)]
        gone = [os.path.basename(path) for path in batch.files if path not in files]
        verb = 'Copied' if batch.copying else 'Opened'
        if not batch.copying and len(files) > MANY_FILES and (not messagebox.askyesno('Open several files', f'Open {len(files)} {what}s?')):
            return
        if files and batch.copying:
            try:
                copy_files_to_clipboard(files, self.winfo_id())
            except OSError as error:
                batch.note(f'Could not copy: {error}')
                files = []
        elif files:
            opened = []
            for path in files:
                try:
                    os.startfile(path)
                    opened.append(path)
                except OSError:
                    gone.append(os.path.basename(path))
            files = opened
        found = len(set().union(*(batch.files[path] for path in files))) if files else 0
        log_activity(f'{verb} {len(files)} {what}s for {rows} rows')
        (text, as_box) = batch_summary(batch.copying, what, rows, found, files, batch.missing + gone, batch.notes)
        self.set_status(text.split('\n')[0])
        if as_box:
            messagebox.showinfo(f'{verb} {what}s', text)

    def _search_clicked_cell(self, event):
        row_id = self.table.identify_row(event.y)
        column_id = self.table.identify_column(event.x)
        if not row_id or not column_id:
            return
        try:
            index = int(column_id.replace('#', '')) - 1
            key = self.visible_columns()[index]
        except (ValueError, IndexError):
            return
        value = self._cell_value(row_id, key)
        if value:
            self.search_var.set(value)
            self.start_search()

    def _cell_value(self, row_id, key):
        source = ROW_SOURCES.get(self.mode, {}).get(key)
        if source is not None and self.results is not None:
            try:
                value = self.results.iloc[int(row_id)].get(source, '')
                text = '' if value is None or value != value else str(value)
                if text.strip():
                    return text.strip()
            except (ValueError, IndexError, KeyError):
                pass
        return str(self.table.set(row_id, key)).strip()

    def _show_context_menu(self, event):
        row_id = self.table.identify_row(event.y)
        if not row_id:
            return
        self._keep_or_select(self.table, row_id)
        self._on_row_selected()
        self._show_menu_states()
        menu = self.component_menu if self.mode == 'components' else self.context_menu
        menu.tk_popup(event.x_root, event.y_root)

    def _revision_numbers(self, number):
        number = str(number).strip()
        siblings = self.index.siblings_sharing_drawing(number)
        parts = dict(self.index.revision_parts(number))
        for sibling in siblings:
            parts.update(self.index.revision_parts(sibling))
        owners = demo_index.drawing_owners(parts)
        hidden = {value for (value, owner) in owners.items() if owner != value and value != number and (not self.index.owns_a_bom(value))}
        for value in hidden:
            parts.pop(value, None)
        return (parts, hidden, siblings)

    def _focus_is_prototype(self):
        if self.components.winfo_ismapped() and self.parts_table.selection():
            return getattr(self, '_components_source', demo_index.PRODUCTION) == demo_index.PROTOTYPE
        row = self.selected_row()
        return row is not None and _is_prototype(row)

    def show_revision_impact(self):
        if self._focus_is_prototype():
            messagebox.showinfo('Released BOMs only', PROTOTYPE_CHECKLIST_NOTE)
            return
        number = self._part_number_in_focus()
        if not number:
            return
        if self.index is None:
            self.set_status('Still loading the index...')
            return
        started = datetime.now()
        (parts, hidden, siblings) = self._revision_numbers(number)
        info = {}
        tree = self.index.revision_tree(number, info=info)
        for sibling in siblings:
            tree.extend(self.index.revision_tree(sibling, info=info))
        elapsed = (datetime.now() - started).total_seconds() * 1000
        log_activity(f'Revision tree for {number}: {len(parts)} parts' + (f' (+{len(hidden)} modifier variants left out)' if hidden else '') + (f' (+{len(siblings)} sharing its drawing)' if siblings else ''))
        if not tree or not any((node['children'] for node in tree)):
            messagebox.showinfo('Nothing else affected', f'{number} is not listed as a part in any BOM, so nothing else would need revising.\n\nIf that is a drawing number, try the full 10-digit part number.')
            return
        if getattr(self, 'revision_panel', None) is not None:
            self.revision_panel.close()
        self.revision_panel = RevisionPanel(self, number, tree, parts, siblings=siblings, truncated=info.get('truncated', False), hidden=hidden)
        self.set_status(f'{len(parts)} part numbers would need revising if {number} changes. Worked out in {elapsed:,.0f} ms.')

    def show_revision_list(self, number=None):
        if number is None and self._focus_is_prototype():
            messagebox.showinfo('Released BOMs only', PROTOTYPE_CHECKLIST_NOTE)
            return
        number = number or self._part_number_in_focus()
        if not number:
            return
        if self.index is None:
            self.set_status('Still loading the index...')
            return
        started = datetime.now()
        matched = self.index.resolve_part_numbers(number)
        expanded = [value for value in matched if value != number]
        (impact, siblings) = self.index.revision_impact_with_siblings(number)
        expanded += siblings
        elapsed = (datetime.now() - started).total_seconds() * 1000
        log_activity(f'Revision impact for {number}: {len(impact)} BOMs')
        if impact.empty:
            messagebox.showinfo('Nothing else affected', f'{number} is not listed as a part in any BOM, so nothing else would need revising.\n\nIf that is a drawing number, try the full 10-digit part number.')
            return
        self.hide_components()
        self.mode_switch.deselect()
        self._mode_changed()
        self.search_var.set(number)
        direct = int((impact['Level'] == 0).sum())
        indirect = len(impact) - direct
        (revving, left_out, _) = self._revision_numbers(number)
        resolved = [value for value in matched if value != number]
        via = ''
        if resolved:
            via += f" (as {', '.join(resolved[:3])}" + ('...' if len(resolved) > 3 else '') + ')'
        named = [value for value in siblings if value not in left_out]
        if named:
            via += f", plus {', '.join(named[:3])}" + ('...' if len(named) > 3 else '') + ' on the same drawing'
        self._show_results(number, impact, elapsed, 'all', caption=f'{len(impact)} BOMs, {len(revving)} part numbers would rev with {number}{via}')
        detail = f'{len(impact)} BOMs and {len(revving)} part numbers would need revising if {number} changes'
        if indirect:
            detail += f" — {direct} BOM{('s' if direct != 1 else '')} list it directly, {indirect} through a sub-assembly that revs with it"
        self.set_status(detail + f'. Found in {elapsed:,.0f} ms.')

    def _part_number_in_focus(self):
        if self.components.winfo_ismapped() and self.parts_table.selection():
            item = self.parts_table.selection()[0]
            return str(self.parts_table.set(item, 'part_number')).strip()
        row = self.selected_row()
        if row is None:
            messagebox.showinfo('Nothing selected', 'Choose a part first, in Components mode or in the parts list of a BOM.')
            return ''
        number = str(row.get('part_number', '')).strip()
        if number:
            return number
        assembly = str(row.get('Internal Assembly No:', '')).strip()
        if not assembly:
            assembly = str(row.get('BOM #', '')).strip().rsplit('-', 1)[0]
        assembly = demo_index.part_number_key(assembly)
        if assembly:
            return assembly
        messagebox.showinfo('No part number on that row', "That row carries no part number to work from. Open the BOM's parts list and choose a part.")
        return ''

    def find_boms_using_part(self):
        row = self.selected_row()
        if row is None:
            return
        number = str(row.get('part_number', '')).strip()
        if not number:
            return
        self.mode_switch.deselect()
        self._mode_changed()
        self.search_var.set(number)
        log_activity(f'Find every BOM using {number}')
        self.start_search()

    def _copy(self, whole_row):
        items = self.table.selection()
        if not items or self.results is None:
            return
        if whole_row:
            lines = ['\t'.join((str(self.table.set(item, key)) for key in self.visible_columns())) for item in items]
        else:
            lines = [str(self._row_number(self.results.iloc[int(item)])) for item in items]
        text = '\n'.join(lines)
        self.clipboard_clear()
        self.clipboard_append(text)
        self.set_status('Copied to clipboard.' if whole_row or len(lines) > 1 else f'Copied {text} to clipboard.')

    def _row_number(self, row):
        if self.mode == 'components':
            return row.get('part_number', '')
        return row.get('BOM #', '')

    def open_bom(self, act=None):
        row = self.selected_row()
        if row is None:
            return
        root = self._folder_for(row, 'bom_root')
        if not root:
            self._no_root(self._row_number(row))
            return
        path = os.path.join(root, row['Folder'], row['Filename'])
        if self.mode == 'components':
            self.set_status(f"Opening {row['Filename']}, one of {row.get('BOMs', 1)} BOMs using {row.get('part_number', '')}.")
        (act or self._open)(path, f"BOM {row['Filename']}")

    def open_drawing(self, act=None):
        row = self.selected_row()
        if row is None:
            return
        self._open_drawing_for(row['Folder'], self._row_number(row), self._folder_for(row, 'drawings_root'), act=act, prototype=_is_prototype(row))

    def show_projects(self):
        row = self.selected_row()
        if row is None:
            return
        assembly = str(row.get('Internal Assembly No:', '')).strip()
        if not assembly:
            assembly = re.sub('-.*$', '', str(row.get('BOM #', '')).strip())
        if getattr(self, 'projects_panel', None) is not None:
            self.projects_panel.close()
        self.projects_panel = ProjectsPanel(self, assembly, row)

    def open_project_folder(self):
        row = self.selected_row()
        if row is None:
            return
        ecn = str(row.get('ECN#:', '')).strip()
        if not ecn:
            messagebox.showinfo('No ECN on this BOM', 'This BOM carries no ECN number, so there is no project to look up.')
            return
        if not self.index.projects:
            messagebox.showinfo('The project table has not been built', 'Project numbers come from the PCR documents. Run\n\n    tools/pcr_projects.py\n\nand the folder listing beside it, then this button will work.')
            return
        project = self.index.project_for(ecn)
        if not project:
            messagebox.showinfo('No project for this ECN', f'ECN {ecn} is not in the project table. Its PCR may be one of the scanned PDFs, which are not read.')
            return
        folders = self.index.project_folders_for(ecn)
        if not folders:
            messagebox.showinfo('No folder for that project', f'ECN {ecn} belongs to project {project}, but no folder starting with that number was found under Active or Closed Projects.')
            return
        folder = folders[0]
        if len(folders) > 1:
            self.set_status(f"Project {project} is in {len(folders)} places; opening the {folder['status']} one.")
        else:
            self.set_status(f"Opening project {project} ({folder['status']}).")
        self._open(folder['path'], f"project folder {folder['name']}")

    def open_customer_drawing(self, act=None):
        row = self.selected_row()
        if row is None:
            return
        if self.mode == 'components':
            self._say('Customer drawings are per assembly', 'Customer prints are filed against an assembly, not an individual part, so there is usually no customer drawing for a component.\n\nSwitch to BOMs and open one of the assemblies this part is used in.', note='Customer drawings are filed per assembly - not looked up for parts.')
            return
        if _is_prototype(row):
            self._say('Not found', 'Customer drawings are not looked up for prototype BOMs.', number=row.get('BOM #', ''))
            return
        commodity = row['Folder']
        folder = os.path.join(demo_config.paths()['drawings_root'], demo_config.drawings_folder_for(commodity), demo_config.customer_print_folder_for(commodity))
        self._open_matching(folder, row.get('BOM #', ''), 'customer drawing', ('.pdf',), act=act)

    def open_dfmea(self, act=None):
        row = self.selected_row()
        if row is None:
            return
        if self.index.dfmea_state() != 'ready':
            self._say(*DFMEA_NOT_BUILT, note='The DFMEA list has not been built.')
            return
        if _is_prototype(row):
            self._say(*DFMEA_NONE, number=row.get('BOM #', ''))
            return
        found = self.index.dfmea_for_bom(row.get('Internal Assembly No:', ''), row.get('BOM #', ''))
        if not found:
            self._say(*DFMEA_NONE, number=row.get('BOM #', ''))
            return
        path = os.path.join(demo_config.paths()['dfmea_root'], found['folder'], found['filename'])
        (act or self._open)(path, f"DFMEA {found['filename']}")

    def open_cad(self, act=None):
        row = self.selected_row()
        if row is None:
            return
        self._open_cad_for(row['Folder'], self._row_number(row), self._folder_for(row, 'math_root'), act=act, prototype=_is_prototype(row))

    def commodities_for_part(self, number):
        number = str(number).strip()
        if not number:
            return []
        found = []
        if self.index is not None:
            components = self.index.components()
            if not components.empty:
                exact = components[components['part_number'].astype(str) == number]
                found = list(dict.fromkeys(exact['Folder'].astype(str)))
        if not found and len(number) >= 2 and number[:2].isdigit():
            prefix = number[:2]
            found = [folder for folder in demo_config.FOLDERS if folder.startswith(prefix)]
        return found

    def quick_open(self, what):
        number = self.quick_var.get().strip()
        if not number:
            self.set_status('Type a part number first.')
            self.quick_box.focus_set()
            return
        if demo_index.prototype_key(number) is not None:
            log_activity(f'Quick open {what} for {number}')
            self._open_prototype(number, 'drawing' if what == 'drawing' else 'cad')
            return
        commodities = self.commodities_for_part(number)
        if not commodities:
            messagebox.showinfo('Which commodity?', f'“{number}” is not in the index, and its first two digits do not match a commodity folder.\n\nSearch for it instead, then use the buttons above.')
            return
        log_activity(f'Quick open {what} for {number}')
        for (position, commodity) in enumerate(commodities):
            last = position == len(commodities) - 1
            folder = os.path.join(demo_config.paths()['drawings_root' if what == 'drawing' else 'math_root'], demo_config.drawings_folder_for(commodity))
            if self._open_matching(folder, number, 'drawing' if what == 'drawing' else 'CAD file', ('.pdf',) if what == 'drawing' else None, quiet=not last):
                return

    def _open_drawing_for(self, commodity, number, root=None, act=None, prototype=False):
        if prototype and needs_prototype_rule(number):
            self._open_prototype(number, 'drawing', act)
            return
        root = demo_config.paths()['drawings_root'] if root is None else root
        if not root:
            self._no_root(number)
            return
        folder = os.path.join(root, demo_config.drawings_folder_for(commodity))
        self._open_matching(folder, number, 'drawing', ('.pdf',), act=act)

    def _open_cad_for(self, commodity, number, root=None, act=None, prototype=False):
        if prototype and needs_prototype_rule(number):
            self._open_prototype(number, 'cad', act)
            return
        root = demo_config.paths()['math_root'] if root is None else root
        if not root:
            self._no_root(number)
            return
        folder = os.path.join(root, demo_config.drawings_folder_for(commodity))
        self._open_matching(folder, number, 'CAD file', None, act=act)

    def _no_root(self, number):
        self.set_status(NO_PROTOTYPE_ROOT)
        batch = getattr(self, '_batch', None)
        if batch is not None:
            batch.note(NO_PROTOTYPE_ROOT)
            self._missed(number)

    def _prototype_paths(self, number, kind):
        root = demo_config.paths().get('proto_drawings_root' if kind == 'drawing' else 'proto_math_root', '')
        if not root:
            self._no_root(number)
            return None
        prototype = getattr(self.index, 'prototype', None)
        if prototype is None:
            self.set_status(PROTOTYPE_NOT_INDEXED)
            self._missed(number)
            return None
        return [os.path.join(root, path) for path in prototype.prototype_files(kind, number)]

    def _open_prototype(self, number, kind, act=None):
        what = 'drawing' if kind == 'drawing' else 'CAD file'
        found = self._prototype_paths(number, kind)
        if found is None:
            return
        act = act or self._open
        if not found:
            self._say('Not found', f'No {what} for {number} in the prototype folders. Obsolete folders are not searched.', number=number)
        elif len(found) == 1:
            act(found[0], f'{what} {os.path.basename(found[0])}')
        elif getattr(self, '_batch', None) is not None:
            for path in found:
                act(path, f'{what} {os.path.basename(path)}')
        else:
            self._pick_file(found, what, act)

    def _pick_file(self, paths, what, act):
        menu = tk.Menu(self, tearoff=0)
        for path in paths:
            name = os.path.basename(path)
            menu.add_command(label=name, command=lambda p=path, n=name: act(p, f'{what} {n}'))
        (x, y) = self.winfo_pointerxy()
        widget = self.focus_get()
        if getattr(self, '_last_input', '') == 'key' and widget is not None:
            (x, y) = (widget.winfo_rootx(), widget.winfo_rooty() + widget.winfo_height())
        menu.tk_popup(x, y)

    def _list_folder(self, folder, extensions):
        if getattr(self, '_folder_cache', None) is None:
            self._folder_cache = {}
        key = (folder, extensions)
        if key not in self._folder_cache:
            try:
                self._folder_cache[key] = [name for name in os.listdir(folder) if not os.path.isdir(os.path.join(folder, name)) and (extensions is None or name.lower().endswith(extensions))]
            except OSError:
                self._folder_cache[key] = []
        return self._folder_cache[key]

    def _find_matching(self, folder, number, extensions, names=None):
        digits = re.match('\\d+', str(number).strip())
        digits = digits.group(0) if digits else ''
        if not digits:
            return ''
        if names is None:
            names = self._list_folder(folder, extensions)
        for wanted in (digits, digits[:demo_index.DRAWING_NUMBER_LENGTH]):
            for name in sorted(names):
                if name.startswith(wanted) or wanted in name:
                    return os.path.join(folder, name)
        return ''

    def _open_matching(self, folder, number, what, extensions, quiet=False, act=None):
        digits = re.match('\\d+', str(number).strip())
        digits = digits.group(0) if digits else ''
        if not digits:
            if not quiet:
                self._say('Nothing to open', f'This row has no number to find a {what} by.', number=number)
            return False
        batch = getattr(self, '_batch', None)
        try:
            names = batch.listing(folder, extensions) if batch is not None else folder_files(folder, extensions)
        except OSError as error:
            if not quiet:
                self._say(f'{what.capitalize()} unavailable', f'Could not read:\n\n{folder}\n\nCheck your network drives.\n\n{error}', error=True, note=f'Could not read {folder} - check your network drives.')
            return False
        found = self._find_matching(folder, number, extensions, names=names)
        if found:
            (act or self._open)(found, f'{what} {os.path.basename(found)}')
            return True
        if not quiet:
            self._say('Not found', f'No {what} for {digits} in:\n\n{folder}', number=digits)
        return False
    SAVE_KINDS = {'.xlsx': 'Excel workbook', '.docx': 'Word document'}

    def ask_where_to_save(self, path):
        extension = os.path.splitext(path)[1].lower() or '.xlsx'
        kind = self.SAVE_KINDS.get(extension, 'File')
        chosen = filedialog.asksaveasfilename(parent=self, title='Save as', initialdir=os.path.dirname(path), initialfile=os.path.basename(path), defaultextension=extension, filetypes=[(kind, '*' + extension)])
        return chosen or ''

    def _copy_file(self, path, description):
        if not os.path.exists(path):
            messagebox.showerror('Cannot copy', missing_file_message(path))
            return
        try:
            copy_files_to_clipboard([path], self.winfo_id())
        except OSError as error:
            messagebox.showerror('Could not copy', f'{path}\n\n{error}')
            return
        log_activity(f'Copied {description}')
        self.set_status(f'Copied {os.path.basename(path)} - paste it into a folder or an email.')

    def _open(self, path, description):
        if not os.path.exists(path):
            messagebox.showerror('Cannot open', missing_file_message(path))
            return
        log_activity(f'Opened {description}')
        try:
            os.startfile(path)
        except OSError as error:
            messagebox.showerror('Could not open', f'{path}\n\n{error}')

    def visible_results(self):
        if self.results is None or self.results.empty:
            return (self.results, 'All')
        chosen = self._tab_labels.get(self.commodity_tabs.get(), 'All')
        if chosen == 'All':
            return (self.results, chosen)
        shown = [_row_is_shown(row, chosen) for (_, row) in self.results.iterrows()]
        return (self.results[shown], chosen)

    def export_results(self):
        if self.results is None or self.results.empty:
            return
        (rows, chosen) = self.visible_results()
        if rows is None or rows.empty:
            return
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        term = self.search_var.get().strip() or 'search'
        safe = ''.join((c if c.isalnum() or c in ' -_' else '_' for c in term)).strip()
        path = os.path.join(default_save_folder(), f'BOM Search Demo - {safe}.xlsx')
        book = Workbook()
        sheet = book.active
        sheet.title = 'Results'
        components = self.mode == 'components'
        sheet.title = 'Parts' if components else 'Results'
        sheet['A1'] = f'BOM Search Demo results for "{term}"'
        sheet['A1'].font = Font(bold=True, size=14)
        scope = 'every commodity' if chosen == 'All' else f'{chosen} only'
        what = 'parts' if components else 'BOMs'
        sheet['A2'] = f'{len(rows):,} {what}, {scope} · {self.index.describe()} · exported {datetime.now():%Y-%m-%d %H:%M}'
        sheet['A2'].font = Font(italic=True, color='666666')
        prototype_fill = None
        if 'Source' in rows.columns and (rows['Source'] == demo_index.PROTOTYPE).any():
            prototype_fill = PatternFill('solid', fgColor=PROTOTYPE_EXPORT_FILL)
            sheet['A3'] = 'Rows shaded pink are prototype BOMs, not released ones.'
            sheet['A3'].font = Font(italic=True, color='666666')
        every_key = [key for (key, *_) in MODE_COLUMNS[self.mode]]
        shown = self.visible_columns()
        headings = [heading for (key, heading, *_) in MODE_COLUMNS[self.mode] if key in shown]
        for (column, heading) in enumerate(headings, start=1):
            cell = sheet.cell(row=4, column=column, value=heading)
            cell.font = Font(bold=True)
        keys = shown
        self._folder_cache = {}
        drawing_column = keys.index('Drawing') + 1 if 'Drawing' in keys else 0
        customer_column = keys.index('CustomerNo') + 1 if 'CustomerNo' in keys else 0
        linked = 0

        def link_cell(column, path):
            if not column or not path:
                return 0
            cell = sheet.cell(row=excel_row, column=column)
            cell.hyperlink = path
            cell.style = 'Hyperlink'
            return 1
        for (offset, (_, row)) in enumerate(rows.iterrows()):
            excel_row = 5 + offset
            values = self._row_values(row)
            values = [values[every_key.index(key)] for key in keys]
            for (column, value) in enumerate(values, start=1):
                sheet.cell(row=excel_row, column=column, value=value)
            root = self._folder_for(row, 'bom_root')
            if root:
                cell = sheet.cell(row=excel_row, column=1)
                cell.hyperlink = os.path.join(root, str(row.get('Folder', '')), str(row.get('Filename', '')))
                cell.style = 'Hyperlink'
            commodity = str(row.get('Folder', ''))
            drawings_root = self._folder_for(row, 'drawings_root')
            if not commodity or not drawings_root:
                continue
            drawings = os.path.join(drawings_root, demo_config.drawings_folder_for(commodity))
            number = self._row_number(row)
            if _is_prototype(row) and needs_prototype_rule(number):
                found = self._prototype_paths(number, 'drawing') or []
                linked += link_cell(drawing_column, found[0] if len(found) == 1 else '')
                continue
            linked += link_cell(drawing_column, self._find_matching(drawings, number, ('.pdf',)))
            if customer_column and (not _is_prototype(row)):
                prints = os.path.join(drawings, demo_config.customer_print_folder_for(commodity))
                linked += link_cell(customer_column, self._find_matching(prints, number, ('.pdf',)))
        self._folder_cache = {}
        if prototype_fill is not None:
            for (offset, (_, row)) in enumerate(rows.iterrows()):
                if _is_prototype(row):
                    for column in range(1, len(keys) + 1):
                        sheet.cell(row=5 + offset, column=column).fill = prototype_fill
        for (column, heading) in enumerate(headings, start=1):
            longest = max([len(str(heading))] + [len(str(sheet.cell(row=r, column=column).value or '')) for r in range(5, 5 + min(len(rows), 200))])
            sheet.column_dimensions[get_column_letter(column)].width = min(longest + 3, 60)
        sheet.freeze_panes = 'A5'
        path = self.ask_where_to_save(path)
        if not path:
            self.set_status('Export cancelled - nothing was written.')
            return
        try:
            book.save(path)
        except OSError as error:
            messagebox.showerror('Could not save', f'{path}\n\n{error}\n\nIs it already open in Excel?')
            return
        log_activity(f'Exported {len(rows)} results ({scope})')
        self.set_status(f'Exported {len(rows):,} results ({scope}) to {path}')
        try:
            os.startfile(path)
        except OSError:
            pass

    def send_feedback(self):
        if getattr(self, 'feedback_panel', None) is not None:
            return
        card = ctk.CTkFrame(self, corner_radius=12, width=660, height=430, border_width=2, border_color=MODE_ACCENTS[self.mode]['normal'])
        card.place(relx=0.5, rely=0.5, anchor='center')
        card.grid_propagate(False)
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(2, weight=1)
        card.lift()
        self.feedback_panel = card

        def close():
            card.destroy()
            self.feedback_panel = None
        ctk.CTkLabel(card, text='What would you like to tell us?', font=ctk.CTkFont(family=FONT, size=16, weight='bold')).grid(row=0, column=0, padx=20, pady=(20, 2), sticky='w')
        ctk.CTkLabel(card, text='Your name and your last search are included automatically.', font=ctk.CTkFont(family=FONT, size=11), text_color=('gray45', 'gray60')).grid(row=1, column=0, padx=20, pady=(0, 10), sticky='w')
        kind = ctk.CTkSegmentedButton(card, values=['Idea', 'Problem', 'Wrong results'], font=ctk.CTkFont(family=FONT, size=12))
        kind.set('Idea')
        kind.grid(row=1, column=0, padx=20, pady=(0, 8), sticky='e')
        box = ctk.CTkTextbox(card, font=ctk.CTkFont(family=FONT, size=13))
        box.grid(row=2, column=0, padx=20, sticky='nsew')
        box.focus_set()
        note = ctk.CTkLabel(card, text='', font=ctk.CTkFont(family=FONT, size=11), text_color=('#a4262c', '#f1707a'))
        note.grid(row=3, column=0, padx=20, sticky='w')
        card.bind('<Escape>', lambda event: close())
        box.bind('<Escape>', lambda event: close())

        def submit():
            message = box.get('1.0', 'end').strip()
            if not message:
                note.configure(text='Type something first.')
                box.focus_set()
                return
            entry = f"{'-' * 70}\n{kind.get()} from {getuser()} at {datetime.now():%Y-%m-%d %H:%M:%S} ({VERSION})\nLast search: {self.search_var.get() or '(none)'}\n\n{message}\n"
            targets = demo_config.feedback_targets()
            (written, failures) = (None, [])
            for target in targets:
                try:
                    folder = os.path.dirname(target)
                    if folder:
                        os.makedirs(folder, exist_ok=True)
                    with open(target, 'a', encoding='utf-8') as handle:
                        handle.write(entry)
                    written = target
                    break
                except OSError as error:
                    failures.append(f'{target}\n  {error}')
            if written is None:
                note.configure(text='Could not save your note - the network folder could not be written. If you are working from home, connect to the VPN and press Send again. Your text is still here.')
                log_activity('Feedback could not be saved: ' + failures[0].replace('\n', ' '))
                return
            log_activity(f'Feedback ({kind.get()}) -> {written}: {message.splitlines()[0][:80]}')
            close()
            self.set_status('Thank you - your note was sent.')
        buttons = ctk.CTkFrame(card, fg_color='transparent')
        buttons.grid(row=4, column=0, padx=20, pady=16, sticky='ew')
        buttons.grid_columnconfigure(0, weight=1)
        ctk.CTkButton(buttons, text='Send', width=120, command=submit, fg_color=MODE_ACCENTS[self.mode]['normal'], hover_color=MODE_ACCENTS[self.mode]['hover'], font=ctk.CTkFont(family=FONT, size=13)).grid(row=0, column=1, padx=6)
        ctk.CTkButton(buttons, text='Cancel', width=100, command=close, fg_color='transparent', border_width=1, text_color=('gray20', 'gray85'), font=ctk.CTkFont(family=FONT, size=13)).grid(row=0, column=2)

    def show_components(self):
        row = self.selected_row()
        if row is None:
            return
        parts = self._index_for(row).parts(row['Filename'])
        if parts.empty:
            messagebox.showinfo('No parts listed', f"No parts table could be read from {row['Filename']}.")
            return
        log_activity(f"Viewed components of {row['Filename']}")
        self._components_commodity = row['Folder']
        self._components_source = row.get('Source', demo_index.PRODUCTION)
        self.parts_revision_button.configure(state='disabled' if self._components_source == demo_index.PROTOTYPE else 'normal')
        self._parts = parts
        self._parts_of = os.path.splitext(row['Filename'])[0]
        self._parts_sort = {}
        self.parts_filter_var.set('')
        useful = [key for (key, _) in demo_index.PART_COLUMNS if parts[key].astype(str).str.strip().any()]
        self.parts_table.configure(displaycolumns=useful or [key for (key, _) in demo_index.PART_COLUMNS])
        self._render_parts()
        self.grid_rowconfigure(5, weight=1)
        self.components.grid()

    def _render_parts(self):
        parts = getattr(self, '_parts', None)
        if parts is None:
            return
        self.parts_table.delete(*self.parts_table.get_children())
        term = self.parts_filter_var.get().strip().lower()
        shown = 0
        thickness_columns = ('thickness_min', 'thickness_max')
        for (_, part) in parts.iterrows():
            values = tuple((_thickness(part[key]) if key in thickness_columns else str(part[key]) for (key, _) in demo_index.PART_COLUMNS))
            if term and (not any((term in value.lower() for value in values))):
                continue
            self.parts_table.insert('', 'end', values=values, tags=('odd' if shown % 2 else 'even',))
            shown += 1
        where = self._parts_of
        if term:
            self.components_label.configure(text=f'{shown} of {len(parts)} parts in {where} matching "{term}"')
        else:
            self.components_label.configure(text=f'{len(parts)} parts in {where}')

    def sort_parts_by(self, key):
        if getattr(self, '_parts', None) is None:
            return
        descending = not self._parts_sort.get(key, False)
        self._parts_sort = {key: descending}
        for (column, heading) in demo_index.PART_COLUMNS:
            arrow = ('  ▾' if descending else '  ▴') if column == key else ''
            try:
                self.parts_table.heading(column, text=heading + arrow)
            except tk.TclError:
                pass
        rows = [(self.parts_table.set(item, key), item) for item in self.parts_table.get_children('')]
        _sort_rows(rows, descending)
        for (position, (_, item)) in enumerate(rows):
            self.parts_table.move(item, '', position)
            self.parts_table.item(item, tags=('odd' if position % 2 else 'even',))

    def hide_components(self):
        self.grid_rowconfigure(5, weight=0)
        self.components.grid_remove()

    def _for_parts(self, opener, act=None, what='file'):
        self._many(self.parts_table, list(self.parts_table.selection()), opener, act, what)

    def _on_part_selected(self, event=None):
        several = len(self.parts_table.selection()) > 1
        prototype = getattr(self, '_components_source', None) == demo_index.PROTOTYPE
        self.parts_where_button.configure(state='disabled' if several else 'normal')
        self.parts_revision_button.configure(state='disabled' if several or prototype else 'normal')
        for label in ('Where else is it used?', 'What if it revs?'):
            off = several or (prototype and label == 'What if it revs?')
            self.parts_menu.entryconfigure(label, state='disabled' if off else 'normal')
        if several:
            self.set_status(f'{len(self.parts_table.selection())} parts selected.')

    def _selected_part(self):
        batch = getattr(self, '_batch', None)
        if batch is not None and batch.owner is self.parts_table:
            return self.parts_table.item(batch.item, 'values')[1]
        selection = self.parts_table.selection()
        if not selection:
            messagebox.showinfo('No part selected', 'Choose a part from the list first.')
            return None
        return self.parts_table.item(selection[0], 'values')[1]

    def open_part_drawing(self, act=None):
        part = self._selected_part()
        if part:
            source = {'Source': getattr(self, '_components_source', demo_index.PRODUCTION)}
            self._open_drawing_for(getattr(self, '_components_commodity', ''), part, self._folder_for(source, 'drawings_root'), act=act, prototype=source['Source'] == demo_index.PROTOTYPE)

    def open_part_cad(self, act=None):
        part = self._selected_part()
        if part:
            source = {'Source': getattr(self, '_components_source', demo_index.PRODUCTION)}
            self._open_cad_for(getattr(self, '_components_commodity', ''), part, self._folder_for(source, 'math_root'), act=act, prototype=source['Source'] == demo_index.PROTOTYPE)

    def copy_part_number(self):
        parts = [self.parts_table.item(item, 'values')[1] for item in self.parts_table.selection()]
        if not parts:
            self._selected_part()
            return
        self.clipboard_clear()
        self.clipboard_append('\n'.join(parts))
        self.set_status(f'Copied {parts[0]} to clipboard.' if len(parts) == 1 else f'Copied {len(parts)} part numbers to clipboard.')

    def _show_parts_menu(self, event):
        row = self.parts_table.identify_row(event.y)
        if not row:
            return
        self._keep_or_select(self.parts_table, row)
        self._on_part_selected()
        self.parts_menu.tk_popup(event.x_root, event.y_root)

    def search_selected_part(self):
        part = self._selected_part()
        if part:
            self.hide_components()
            self.search_var.set(part)
            self.start_search()

    def _on_escape(self, event=None):
        if getattr(self, 'projects_panel', None) is not None:
            self.projects_panel.close()
        elif getattr(self, 'revision_panel', None) is not None:
            self.revision_panel.close()
        elif self.components.winfo_ismapped():
            self.hide_components()
        else:
            self.search_var.set('')
            self._focus_search()

    def _show_help(self):
        document = demo_config.help_document()
        if document:
            log_activity('Opened the work instruction')
            try:
                webbrowser.open('file:///' + document.replace('\\', '/'))
                return
            except Exception:
                pass
        lines = [f'{VERSION}', '']
        for (release, changes) in CHANGELOG:
            lines.append(f'Version {release}')
            lines.extend((f'   • {change}' for change in changes))
        lines += ['', 'Type anything - a drawing number, a part name, a customer code.', 'Select a result to open its BOM, drawing or CAD file.', '', f"Index: {(self.index.describe() if self.index else 'not loaded')}"]
        messagebox.showinfo(f'About {VERSION}', '\n'.join(lines))

    def _on_close(self):
        save_settings(self.settings)
        self.destroy()

class ProjectsPanel(ctk.CTkFrame):

    def __init__(self, app, assembly, row):
        super().__init__(app, corner_radius=12, width=820, height=560, border_width=2, border_color=MODE_ACCENTS[app.mode]['normal'])
        self.app = app
        self.place(relx=0.5, rely=0.5, anchor='center')
        self.grid_propagate(False)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)
        self.lift()
        header = ctk.CTkFrame(self, fg_color='transparent')
        header.grid(row=0, column=0, sticky='ew', padx=20, pady=(16, 0))
        header.grid_columnconfigure(0, weight=1)
        name = _one_line(row.get('Assembly Name:', ''))
        ctk.CTkLabel(header, text='Projects for %s' % (assembly or 'this BOM'), font=ctk.CTkFont(family=FONT, size=18, weight='bold'), anchor='w').grid(row=0, column=0, sticky='ew')
        ctk.CTkButton(header, text='Close', width=76, height=28, font=ctk.CTkFont(family=FONT, size=12), command=self.close).grid(row=0, column=1, sticky='e')
        purpose = 'Every project this part has been released under.'
        ctk.CTkLabel(self, text='%s  ·  %s' % (name, purpose) if name else purpose, font=ctk.CTkFont(family=FONT, size=12), text_color=('gray40', 'gray65'), anchor='w').grid(row=1, column=0, sticky='ew', padx=20, pady=(2, 10))
        self.list = ctk.CTkScrollableFrame(self, fg_color=('gray94', 'gray17'))
        self.list.grid(row=2, column=0, sticky='nsew', padx=16, pady=(0, 8))
        self.list.grid_columnconfigure(0, weight=1)
        self.footer = ctk.CTkLabel(self, text='', font=ctk.CTkFont(family=FONT, size=11), text_color=('gray45', 'gray60'), anchor='w', justify='left', wraplength=760)
        self.footer.grid(row=3, column=0, sticky='ew', padx=20, pady=(0, 14))
        self._fill(assembly, row)
        self.bind('<Escape>', lambda event: self.close())
        self.after(60, self.lift)

    def close(self):
        self.destroy()
        self.app.projects_panel = None

    def _fill(self, assembly, row):
        index = self.app.index
        entries = index.projects_for_assembly(assembly) if index and assembly else []
        if not entries:
            self._explain(index, row)
            return
        (row_number, heading_shown) = (0, False)
        for entry in entries:
            if not entry['current'] and (not heading_shown):
                self._heading(row_number, 'Obsolete project folders', 'Every project the earlier revisions of this part were released under, newest first.' if row_number else 'No BOM in production carries a project. These are from the earlier revisions, newest first.')
                heading_shown = True
                row_number += 1
            self._add(row_number, entry)
            row_number += 1
        obsolete_read = index is not None and getattr(index.obsolete, 'empty', True) is False
        older = sum((1 for entry in entries if not entry['current']))
        if obsolete_read:
            text = '%d project(s): the current one from the BOM in production, and %d from its obsolete revisions.' % (len(entries), older)
        else:
            text = 'Only the BOM in production has been read. The obsolete BOMs are not indexed yet, so any earlier project is missing from this list rather than absent from the archive.'
        sampled = sorted(getattr(index, 'sampled', ()) or ())
        if sampled:
            which = ' and '.join(sampled)
            noun = 'table is a SAMPLE' if len(sampled) == 1 else 'tables are SAMPLES'
            text += '   |   INCOMPLETE: the ' + which + ' ' + noun + ', not the whole archive. A project missing from this list may simply not be in the sample.'
        if not getattr(index, 'project_folders', None):
            text += '   |   The folder listing has not been built, so no folder can be opened. Run tools/project_folders.py.'
        self.footer.configure(text=text)

    def _heading(self, position, text, note):
        block = ctk.CTkFrame(self.list, fg_color='transparent')
        block.grid(row=position, column=0, sticky='ew', pady=(10, 6), padx=2)
        block.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(block, text=text, anchor='w', font=ctk.CTkFont(family=FONT, size=12, weight='bold'), text_color=('gray35', 'gray70')).grid(row=0, column=0, sticky='w', padx=(12, 0))
        ctk.CTkLabel(block, text=note, anchor='w', justify='left', wraplength=600, font=ctk.CTkFont(family=FONT, size=11), text_color=('gray50', 'gray55')).grid(row=1, column=0, sticky='w', padx=(12, 0))

    def _add(self, position, entry):
        current = entry['current']
        card = ctk.CTkFrame(self.list, corner_radius=8, fg_color=('#e8f0fb', '#22303f') if current else ('gray98', 'gray21'))
        card.grid(row=position, column=0, sticky='ew', pady=(0, 6), padx=2)
        card.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(card, text=entry['project'], font=ctk.CTkFont(family=FONT, size=15, weight='bold'), width=76, anchor='w').grid(row=0, column=0, sticky='w', padx=(14, 8), pady=(10, 0))
        index = getattr(self.app, 'index', None)
        listed = bool(getattr(index, 'project_folders', None))
        ctk.CTkLabel(card, text=entry['name'] or ('(folder not in the listing)' if listed else '(no folder listing has been built)'), font=ctk.CTkFont(family=FONT, size=13), anchor='w').grid(row=0, column=1, sticky='ew', pady=(10, 0))
        ctk.CTkLabel(card, text='IN PRODUCTION' if current else 'obsolete', font=ctk.CTkFont(family=FONT, size=10, weight='bold'), text_color=('#1a6fd4', '#7fb3f0') if current else ('gray50', 'gray55')).grid(row=0, column=2, padx=(8, 14), pady=(10, 0))
        detail = 'from BOM %s' % entry['bom']
        if entry['ecn']:
            detail += '   ·   ECN %s' % entry['ecn']
        ctk.CTkLabel(card, text=detail, font=ctk.CTkFont(family=FONT, size=11), text_color=('gray45', 'gray60'), anchor='w').grid(row=1, column=0, columnspan=2, sticky='w', padx=(14, 8), pady=(0, 10))
        buttons = ctk.CTkFrame(card, fg_color='transparent')
        buttons.grid(row=1, column=2, padx=(8, 14), pady=(0, 10))
        open_button = ctk.CTkButton(buttons, text='Open folder', width=104, height=28, font=ctk.CTkFont(family=FONT, size=12), state='normal' if entry['path'] else 'disabled', command=lambda path=entry['path'], project=entry['project']: self.app._open(path, 'project folder %s' % project))
        open_button.pack(side='left')
        if len(entry['folders']) > 1:
            explain(open_button, 'This project has more than one folder - it was reopened. This opens the active one.')

    def _explain(self, index, row):
        ecn = str(row.get('ECN#:', '')).strip()
        if index is None:
            message = 'The index is not loaded.'
        elif not index.projects:
            message = 'The ECN-to-project table has not been built yet, so no project can be looked up for any BOM.\n\nIt comes from tools/pcr_projects.py, which reads the PCR documents.'
        elif not ecn:
            message = 'This BOM carries no ECN number, so there is no change to trace back to a project.'
        elif not index.project_for(ecn):
            message = 'ECN %s is not in the project table. Its PCR may be one of the scanned PDFs, which are not read.' % ecn
        else:
            message = 'The project number was found, but no folder for it is in the folder listing.\n\nRun tools/project_folders.py to refresh that list.'
        ctk.CTkLabel(self.list, text=message, wraplength=620, justify='left', font=ctk.CTkFont(family=FONT, size=13), anchor='w').grid(row=0, column=0, sticky='w', padx=14, pady=14)

class RevisionPanel(ctk.CTkFrame):
    COLUMNS = [('name', 'Name', 280, 'w'), ('update', 'What to update', 250, 'w'), ('customer', 'Customer / Platform', 165, 'w'), ('where', 'Commodity', 150, 'w'), ('done', 'Done', 60, 'center')]

    def __init__(self, app, number, tree, parts, siblings=(), truncated=False, hidden=()):
        width = max(880, min(1240, app.winfo_width() - 120 or 1240))
        height = max(520, min(760, app.winfo_height() - 120 or 760))
        super().__init__(app, corner_radius=12, width=width, height=height, border_width=2, border_color=MODE_ACCENTS[app.mode]['normal'])
        self.app = app
        self.number = number
        self.tree_data = tree
        self.parts = parts
        self.owners = demo_index.drawing_owners(parts)
        self.hidden = set(hidden)
        self.hidden_owners = demo_index.drawing_owners(list(parts) + list(hidden))
        self.root_rows = {}
        self.siblings = list(siblings)
        self.truncated = truncated
        self.rows = {}
        self.ticked = set()
        self.place(relx=0.5, rely=0.5, anchor='center')
        self.grid_propagate(False)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)
        self._build_header()
        self._build_tree()
        self._build_footer()
        self._populate()
        self.bind('<Escape>', lambda event: self.close())
        self.after(80, self._raise)

    def close(self):
        self.destroy()
        self.app.revision_panel = None

    def _raise(self):
        self.lift()
        self.focus_set()

    def _build_header(self):
        header = ctk.CTkFrame(self, corner_radius=0, fg_color='transparent')
        header.grid(row=0, column=0, sticky='ew', padx=6, pady=(8, 0))
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(header, text=f'What needs revising if {self.number} changes', font=ctk.CTkFont(family=FONT, size=17, weight='bold'), anchor='w').grid(row=0, column=0, sticky='w', padx=16, pady=(12, 0))
        ctk.CTkButton(header, text='Close', width=76, height=28, font=ctk.CTkFont(family=FONT, size=12), command=self.close).grid(row=0, column=1, sticky='e', padx=(8, 16), pady=(12, 0))
        resolved = self.app.index.resolve_part_numbers(self.number)
        first = resolved[0] if resolved else self.number
        name = self.app.index.part_name(first)
        subtitle = f'{self.number}'
        if first != self.number:
            subtitle += f'  (as {first})'
        if name:
            subtitle += f'  ·  {name}'
        ctk.CTkLabel(header, text=subtitle, font=ctk.CTkFont(family=FONT, size=13), text_color=('gray35', 'gray70'), anchor='w').grid(row=1, column=0, columnspan=2, sticky='w', padx=16, pady=(2, 0))
        boms = {node['filename'] for node in self._walk(self.tree_data) if node['kind'] == 'bom'}
        shown = {node['part'] for node in self._walk(self.tree_data) if node['kind'] == 'part'}
        shown |= {node['assembly'] for node in self._walk(self.tree_data) if node['kind'] == 'bom' and node.get('assembly')}
        shown &= set(self.parts)
        summary = f"{len(self.parts)} part numbers need revising, and {len(boms)} BOM document{('s' if len(boms) != 1 else '')} reissued."
        listed = [value for value in self.siblings if value not in self.hidden]
        if listed:
            summary += f"  Includes {', '.join(listed)}, drawn on the same drawing."
        asked = str(self.number).strip()
        asked_owner = self.owners.get(asked, asked)
        if asked_owner != asked:
            summary += f'  {asked} is itself a modifier variant: its drawing and CAD are {asked_owner}’s, which is on this list.'
        if self.hidden:
            one = len(self.hidden) == 1
            summary += f"  {len(self.hidden)} more {('differs' if one else 'differ')} only by the modifier digit, so {('it shares' if one else 'they share')} a drawing and CAD with a number on this list and {('is not listed itself' if one else 'are not listed themselves')}."
        covered = sum((1 for (number, owner) in self.owners.items() if owner != number and self._owns_a_bom(number)))
        if covered:
            summary += f'  {covered} shown here share another number’s drawing and CAD but have a BOM of their own to reissue.'
        if self.app.index.dfmea_state() != 'ready':
            summary += '  ' + DFMEA_NOT_NAMED
        elif any((self.app.index.dfmea_for(number) for number in self.parts if self._owns_a_bom(number))):
            summary += "  A DFMEA keeps its own revision, independent of its BOM's, so the two letters often differ."
        if self.truncated:
            summary += '  THE LIST IS CUT SHORT: too many branches to expand, so some numbers are missing. Export it and tell me.'
        elif len(shown) < len(self.parts):
            summary += f'  Only {len(shown)} of them appear as rows here - Export checklist lists all {len(self.parts)}.'
        ctk.CTkLabel(header, text=summary, font=ctk.CTkFont(family=FONT, size=12), text_color=('gray35', 'gray70'), anchor='w', justify='left', wraplength=1100).grid(row=2, column=0, columnspan=2, sticky='w', padx=16, pady=(2, 12))
        buttons = ctk.CTkFrame(self, corner_radius=0, fg_color='transparent')
        buttons.grid(row=1, column=0, sticky='ew', padx=12, pady=(10, 6))
        for (label, command) in (('Expand all', self.expand_all), ('Collapse all', self.collapse_all), ('Copy as text', self.copy_text), ('Export checklist', self.export), ('Draft ECN', self.export_ecn)):
            ctk.CTkButton(buttons, text=label, width=124, height=32, font=ctk.CTkFont(family=FONT, size=12), command=command).pack(side='left', padx=(0, 8))
        ctk.CTkButton(buttons, text='Show as a flat list', width=150, height=32, font=ctk.CTkFont(family=FONT, size=12), fg_color=('gray75', 'gray30'), hover_color=('gray68', 'gray38'), command=self.show_flat).pack(side='right')

    def _build_tree(self):
        holder = tk.Frame(self, borderwidth=0, highlightthickness=0)
        holder.grid(row=2, column=0, sticky='nsew', padx=12, pady=(0, 6))
        holder.grid_columnconfigure(0, weight=1)
        holder.grid_rowconfigure(0, weight=1)
        self.tree = ttk.Treeview(holder, columns=[key for (key, *_) in self.COLUMNS], style='Demo.Treeview', selectmode='extended')
        self.tree.heading('#0', text='Part number  /  BOM', anchor='w')
        self.tree.column('#0', width=380, minwidth=220, anchor='w', stretch=False)
        for (key, heading, width, anchor) in self.COLUMNS:
            self.tree.heading(key, text=heading, anchor=anchor)
            self.tree.column(key, width=width, anchor=anchor, stretch=key == 'name')
        self.tree.grid(row=0, column=0, sticky='nsew')
        bar = ttk.Scrollbar(holder, orient='vertical', command=self.tree.yview)
        bar.grid(row=0, column=1, sticky='ns')
        sideways = ttk.Scrollbar(holder, orient='horizontal', command=self.tree.xview)
        sideways.grid(row=1, column=0, sticky='ew')
        self.tree.configure(yscrollcommand=bar.set, xscrollcommand=sideways.set)
        self.tree.bind('<Button-1>', self._clicked)
        self.tree.bind('<space>', lambda event: self._toggle(self.tree.focus()))
        self.tree.bind('<Double-1>', lambda event: self._open_default())
        for sequence in ('<Control-a>', '<Control-A>'):
            self.tree.bind(sequence, lambda event: self.app._select_all(self.tree))
        self.tree.bind('<<TreeviewSelect>>', lambda event: self._on_rows_selected())
        copy = self.app._copy_file
        self.menu = tk.Menu(self, tearoff=0)
        openers = ((self.open_bom, 'BOM'), (self.open_drawing, 'drawing'), (self.open_cad, 'CAD file'), (self.open_dfmea, 'DFMEA'))
        for (label, (opener, what)) in zip(('Open BOM', 'Open drawing', 'Open CAD', 'Open DFMEA'), openers):
            self.menu.add_command(label=label, command=lambda o=opener, w=what: self._for_rows(o, what=w))
        self.menu.add_separator()
        for (label, (opener, what)) in zip(('Copy BOM file', 'Copy drawing file', 'Copy CAD file', 'Copy DFMEA file'), openers):
            command = lambda o=opener, w=what: self._for_rows(o, act=copy, what=w)
            self.menu.add_command(label=label, command=command)
        self.menu.add_separator()
        self.menu.add_command(label='Copy number', command=self.copy_number)
        self.tree.bind('<Button-3>', self._show_menu)

    def _build_footer(self):
        footer = ctk.CTkFrame(self, corner_radius=0, fg_color='transparent')
        footer.grid(row=3, column=0, sticky='ew', padx=12, pady=(0, 12))
        ctk.CTkLabel(footer, text='Selected row:', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray35', 'gray70')).pack(side='left', padx=(0, 10))
        for (label, command, what) in (('Open BOM', self.open_bom, 'BOM'), ('Drawing', self.open_drawing, 'drawing'), ('CAD', self.open_cad, 'CAD file'), ('DFMEA', self.open_dfmea, 'DFMEA')):
            command = lambda o=command, w=what: self._for_rows(o, what=w)
            ctk.CTkButton(footer, text=label, width=118, height=32, font=ctk.CTkFont(family=FONT, size=12), command=command).pack(side='left', padx=(0, 8))
        self.status = ctk.CTkLabel(footer, text='Tick a row as you finish it.', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray35', 'gray70'))
        self.status.pack(side='right')

    def _owns_a_bom(self, part):
        return self.app.index.owns_a_bom(part)

    def _populate(self):
        for node in self.tree_data:
            commodity = ''
            for candidate in self.app.commodities_for_part(node['part']):
                commodity = candidate
                break
            self._insert('', node, self.app.index.files, commodity)
        self.expand_all()
        first = self.tree.get_children()
        if first:
            self.tree.selection_set(first[0])
            self.tree.focus(first[0])

    def _what_to_update(self, number):
        owner = self.owners.get(number, number)
        covered = owner != number
        dfmea = self.app.index.dfmea_for(number) if self._owns_a_bom(number) else None
        return update_text(self._owns_a_bom(number), covered, owner, dfmea)

    def _insert(self, parent, node, files, commodity=''):
        if node['kind'] == 'part' and node['part'] in self.hidden:
            if parent == '':
                parent = self.root_rows.get(self.hidden_owners.get(node['part']), '')
            for child in node['children']:
                self._insert(parent, child, files, commodity)
            return parent
        if node['kind'] == 'part':
            what = self._what_to_update(node['part'])
            label = node['part']
            if node['repeat']:
                label += '   (listed above)'
            values = (node['name'], what, '', demo_config.commodity_name(commodity), UNTICKED)
        else:
            commodity = node.get('Folder', '')
            assembly = node.get('assembly', '')
            label = node['filename']
            what = 'the BOM document'
            if assembly:
                label = f"{assembly}   ({node['filename']})"
                if node.get('repeat'):
                    label += '   (listed above)'
                what = self._what_to_update(assembly)
            values = (node.get('Assembly Name:', ''), what, node.get('Customer & Platform:', ''), demo_config.commodity_name(commodity), UNTICKED)
        item = self.tree.insert(parent, 'end', text=label, values=values, open=True)
        node['commodity'] = commodity
        self.rows[item] = node
        if parent == '' and node['kind'] == 'part':
            self.root_rows[node['part']] = item
        for child in node['children']:
            self._insert(item, child, files, commodity)
        return item

    def _walk(self, nodes):
        for node in nodes:
            yield node
            for found in self._walk(node['children']):
                yield found

    def _clicked(self, event):
        if self.tree.identify_region(event.x, event.y) != 'cell':
            return
        if self.tree.identify_column(event.x) == f'#{len(self.COLUMNS)}':
            self._toggle(self.tree.identify_row(event.y))
            return 'break'

    def _toggle(self, item):
        if not item:
            return
        if item in self.ticked:
            self.ticked.discard(item)
            self.tree.set(item, 'done', UNTICKED)
        else:
            self.ticked.add(item)
            self.tree.set(item, 'done', TICKED)
        self.status.configure(text=f'{len(self.ticked)} of {len(self.rows)} rows ticked off.')

    def expand_all(self):
        for item in self.rows:
            self.tree.item(item, open=True)

    def collapse_all(self):
        for item in self.rows:
            self.tree.item(item, open=False)
        tops = self.tree.get_children('')
        if tops:
            self.tree.selection_set(tops[0])
            self.tree.focus(tops[0])
            self.tree.see(tops[0])

    def _for_rows(self, opener, act=None, what='file'):
        self.app._many(self.tree, list(self.tree.selection()), opener, act, what)

    def _on_rows_selected(self):
        count = len(self.tree.selection())
        if count > 1:
            self.status.configure(text=f'{count} rows selected.')

    def _selected(self):
        batch = getattr(self.app, '_batch', None)
        if batch is not None and batch.owner is self.tree:
            return (batch.item, self.rows.get(batch.item))
        selection = self.tree.selection()
        focus = self.tree.focus()
        item = focus if focus in selection else selection[0] if selection else ''
        return (item, self.rows.get(item))

    def _open_default(self):
        (item, node) = self._selected()
        if node and node['kind'] == 'bom':
            self.open_bom()
        else:
            self.open_drawing()

    def open_bom(self, act=None):
        (item, node) = self._selected()
        if not node:
            return
        act = act or self.app._open
        if node['kind'] == 'bom':
            path = os.path.join(demo_config.paths()['bom_root'], node.get('Folder', ''), node['filename'])
            act(path, f"BOM {node['filename']}")
            return
        ancestor = self.tree.parent(item)
        while ancestor:
            node_above = self.rows.get(ancestor)
            if node_above and node_above['kind'] == 'bom':
                path = os.path.join(demo_config.paths()['bom_root'], node_above.get('Folder', ''), node_above['filename'])
                act(path, f"BOM {node_above['filename']}")
                return
            ancestor = self.tree.parent(ancestor)
        self.app._say('No BOM above this row', f"{node['part']} is not inside a BOM in this list, so there is no workbook to open for it.", number=node['part'])

    def open_drawing(self, act=None):
        (item, node) = self._selected()
        if not node:
            return
        number = node['part'] if node['kind'] == 'part' else node.get('assembly') or node.get('BOM #', '')
        commodity = node.get('commodity') or ''
        if not commodity:
            for candidate in self.app.commodities_for_part(number):
                commodity = candidate
                break
        self.app._open_drawing_for(commodity, number, act=act)

    def open_cad(self, act=None):
        (item, node) = self._selected()
        if not node:
            return
        number = node['part'] if node['kind'] == 'part' else node.get('assembly', '')
        if not number:
            self.app._say('CAD is per part', 'This row carries no part number to find CAD by.', number=node.get('filename', ''))
            return
        commodity = node.get('commodity') or ''
        if not commodity:
            for candidate in self.app.commodities_for_part(number):
                commodity = candidate
                break
        self.app._open_cad_for(commodity, number, act=act)

    def open_dfmea(self, act=None):
        (item, node) = self._selected()
        if not node:
            return
        index = self.app.index
        if index.dfmea_state() != 'ready':
            self.app._say(*DFMEA_NOT_BUILT, note='The DFMEA list has not been built.')
            return
        found = None
        if node['kind'] == 'bom':
            found = index.dfmea_for_bom(node.get('assembly'), node.get('BOM #'))
        elif self._owns_a_bom(node['part']):
            found = index.dfmea_for(node['part'])
        else:
            ancestor = self.tree.parent(item)
            while ancestor:
                above = self.rows.get(ancestor) or {}
                if above.get('kind') == 'bom':
                    found = index.dfmea_for_bom(above.get('assembly'), above.get('BOM #'))
                    break
                ancestor = self.tree.parent(ancestor)
        if not found:
            self.app._say(*DFMEA_NONE, number=node.get('part') or node.get('assembly', ''))
            return
        path = os.path.join(demo_config.paths()['dfmea_root'], found['folder'], found['filename'])
        (act or self.app._open)(path, f"DFMEA {found['filename']}")

    def copy_number(self):
        numbers = []
        for item in self.tree.selection():
            node = self.rows.get(item)
            if node:
                numbers.append(node['part'] if node['kind'] == 'part' else node.get('assembly') or node['filename'])
        if not numbers:
            return
        self.clipboard_clear()
        self.clipboard_append('\n'.join(numbers))
        self.status.configure(text=f'Copied {numbers[0]} to clipboard.' if len(numbers) == 1 else f'Copied {len(numbers)} numbers to clipboard.')

    def _show_menu(self, event):
        row = self.tree.identify_row(event.y)
        if not row:
            return
        self.app._keep_or_select(self.tree, row)
        self.menu.tk_popup(event.x_root, event.y_root)

    def _lines(self):
        lines = []

        def walk(item, depth):
            node = self.rows[item]
            tick = '[x]' if item in self.ticked else '[ ]'
            if node['kind'] == 'part':
                what = self.tree.set(item, 'update')
                lines.append('%s%s %s  %s  (%s)' % ('    ' * depth, tick, node['part'], node['name'], what))
            else:
                lines.append('%s%s %s  %s  (%s)' % ('    ' * depth, tick, node.get('assembly') or node['filename'], node.get('Assembly Name:', ''), self.tree.set(item, 'update')))
            for child in self.tree.get_children(item):
                walk(child, depth + 1)
        for root in self.tree.get_children():
            walk(root, 0)
        return lines

    def copy_text(self):
        text = '\n'.join([f'If {self.number} changes, revise:'] + self._lines())
        self.clipboard_clear()
        self.clipboard_append(text)
        self.status.configure(text='Copied to the clipboard.')

    def _artefacts(self):
        index = self.app.index
        (boms, customers, commodities, names) = ({}, {}, {}, {})
        customer_parts = {}
        for node in self._walk(self.tree_data):
            if node['kind'] == 'part':
                if node.get('name'):
                    names.setdefault(node['part'], node['name'])
                continue
            key = node.get('assembly') or ''
            if not key:
                continue
            boms[key] = node['filename']
            customers[key] = node.get('Customer & Platform:', '')
            customer_parts[key] = node.get('Customer Assembly No:', '')
            commodities[key] = node.get('Folder', '')
            if node.get('Assembly Name:'):
                names.setdefault(key, node['Assembly Name:'])
        ticked_numbers = set()
        for item in self.ticked:
            node = self.rows.get(item, {})
            ticked_numbers.add(node.get('part') if node.get('kind') == 'part' else node.get('assembly'))
        ticked_numbers.discard(None)
        found = []
        for number in sorted(self.parts):
            commodity = commodities.get(number, '')
            inferred = False
            if not commodity:
                candidates = self.app.commodities_for_part(number)
                commodity = candidates[0] if candidates else ''
                inferred = bool(commodity)
            drawing = index.drawing_for(number)
            owner = self.owners.get(number, number)
            covered = owner != number
            found.append({'part': number, 'name': names.get(number) or index.part_name(number), 'commodity': commodity, 'inferred': inferred, 'drawing': '' if covered else drawing + '.pdf' if drawing else '', 'step': '' if covered else number + '.stp' if number.isdigit() else '', 'bom': boms.get(number, ''), 'customer': customers.get(number, ''), 'customer_part': customer_parts.get(number, ''), 'covered_by': owner if covered else '', 'done': number in ticked_numbers})
        return found

    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        entries = self._artefacts()
        paths = demo_config.paths()
        safe = ''.join((c if c.isalnum() or c in ' -_' else '_' for c in str(self.number))).strip()
        path = os.path.join(default_save_folder(), f'Demo revision checklist - {safe}.xlsx')
        face = 'Arial'
        head_fill = PatternFill('solid', fgColor='1F2429')
        head_font = Font(name=face, bold=True, color='FFFFFF', size=10)
        title_font = Font(name=face, bold=True, size=14)
        note_font = Font(name=face, italic=True, size=9, color='666666')
        body = Font(name=face, size=10)
        bold = Font(name=face, size=10, bold=True)
        muted = Font(name=face, size=10, color='888888')
        edge = Border(bottom=Side(style='thin', color='D6DBE0'))

        def heading(sheet, headings, widths, subtitle):
            sheet['A1'] = f'Everything to revise if {self.number} changes'
            sheet['A1'].font = title_font
            sheet['A2'] = subtitle
            sheet['A2'].font = note_font
            for (column, (text, width)) in enumerate(zip(headings, widths), start=1):
                cell = sheet.cell(row=4, column=column, value=text)
                cell.font = head_font
                cell.fill = head_fill
                sheet.column_dimensions[get_column_letter(column)].width = width
            sheet.row_dimensions[4].height = 20
            sheet.freeze_panes = 'A5'

        def put(sheet, row, values, fonts=None):
            for (column, value) in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = (fonts or {}).get(column, body)
                cell.border = edge
                cell.alignment = Alignment(vertical='center')
        band = PatternFill('solid', fgColor='E8EDF2')
        band_font = Font(name=face, bold=True, size=11)

        def section(sheet, row, title, width):
            sheet.cell(row=row, column=1, value=title).font = band_font
            for column in range(1, width + 1):
                sheet.cell(row=row, column=column).fill = band
            sheet.row_dimensions[row].height = 18
            return row + 1
        book = Workbook()
        todo = book.active
        todo.title = 'To do'
        (drawings, steps, workbooks) = ([], [], [])
        for entry in entries:
            math_folder = os.path.join(paths['math_root'], demo_config.drawings_folder_for(entry['commodity']))
            if entry['drawing']:
                drawings.append((entry, 'Update drawing', entry['drawing'], os.path.join(paths['drawings_root'], demo_config.drawings_folder_for(entry['commodity']))))
            if entry['step']:
                steps.append((entry, 'Update STEP file', entry['step'], math_folder))
            if entry['bom']:
                workbooks.append((entry, 'Reissue BOM', entry['bom'], os.path.join(paths['bom_root'], entry['commodity'])))
        dfmeas = dfmea_actions(entries, self.app.index.dfmea_for, paths.get('dfmea_root', ''))
        groups = [('Drawings', drawings), ('STEP files', steps), ('BOMs', workbooks), ('DFMEAs', dfmeas)]
        actions = [action for (_, group) in groups for action in group]
        heading(todo, ['Done', 'Part number', 'Name', 'What to do', 'File', 'Folder'], [7, 15, 34, 18, 26, 46], f'{len(actions)} things to do, across {len(entries)} part numbers, grouped by the kind of document. One row per file to open. STEP file names are the name to look for - nothing indexes the Math folder.')
        line = 5
        for (title, group) in groups:
            if not group:
                continue
            line = section(todo, line, f'{title}  ({len(group)})', 6)
            for (entry, what, filename, folder) in group:
                put(todo, line, ['x' if entry['done'] else '', entry['part'], entry['name'], what, filename, folder], fonts={2: bold, 4: bold})
                line += 1
        if workbooks:
            line = section(todo, line, 'Because the BOM changes', 6)
            for (what, required, why) in bom_change_followups(self.app.index.dfmea_state(), len(workbooks)):
                put(todo, line, ['', '', why, what, '', ''], fonts={3: body if required else note_font, 4: bold if required else muted})
                line += 1
        todo.auto_filter.ref = f'A4:F{line - 1}'
        by_part = book.create_sheet('By part number')
        owning = sum((1 for entry in entries if entry['bom']))
        heading(by_part, ['Done', 'Part number', 'Name', 'Commodity', 'Drawing to update', 'STEP file (expected name)', 'BOM to reissue', 'Customer / Platform'], [7, 15, 34, 26, 20, 26, 22, 24], f'{len(entries)} part numbers, {owning} of them with a BOM to reissue. A dash means there is nothing of that kind to update. Where a number differs from another only by its modifier digit - 7001234201 beside 7001234101 - the drawing and the CAD are the same files, and are listed once against the number that owns them.')
        for (offset, entry) in enumerate(entries):
            commodity = demo_config.commodity_name(entry['commodity'])
            if entry['inferred']:
                commodity += '  (from the part number)'
            shared = f"with {entry['covered_by']}" if entry['covered_by'] else '-'
            put(by_part, 5 + offset, ['x' if entry['done'] else '', entry['part'], entry['name'], commodity, entry['drawing'] or shared, entry['step'] or shared, entry['bom'] or '-', entry['customer']], fonts={2: bold, 5: body if entry['drawing'] else muted, 6: body if entry['step'] else muted, 7: body if entry['bom'] else muted})
        by_part.auto_filter.ref = f'A4:H{4 + len(entries)}'
        why = book.create_sheet('Why')
        heading(why, ['Done', 'Level', 'Part number / BOM', 'Name', 'What to update', 'Customer / Platform', 'Commodity'], [7, 7, 40, 34, 22, 24, 20], 'The tree as the app shows it, so you can see how each number was reached. A number reached twice is expanded the first time and marked "(listed above)" after that; the By part number sheet is the complete list.')
        line = [5]

        def descend(item, depth):
            node = self.rows[item]
            is_part = node['kind'] == 'part'
            label = node['part'] if is_part else node.get('assembly') or node['filename']
            put(why, line[0], ['x' if item in self.ticked else '', depth, label, node['name'] if is_part else node.get('Assembly Name:', ''), self.tree.set(item, 'update'), '' if is_part else node.get('Customer & Platform:', ''), demo_config.commodity_name(node.get('commodity', ''))], fonts={3: bold})
            why.cell(row=line[0], column=3).alignment = Alignment(indent=depth, vertical='center')
            line[0] += 1
            for child in self.tree.get_children(item):
                descend(child, depth + 1)
        for root in self.tree.get_children():
            descend(root, 0)
        path = self.app.ask_where_to_save(path)
        if not path:
            self.app.set_status('Export cancelled - nothing was written.')
            return
        try:
            book.save(path)
        except OSError as error:
            messagebox.showerror('Could not save', f'{path}\n\n{error}')
            return
        log_activity(f'Exported revision checklist for {self.number}: {len(actions)} actions, {len(entries)} part numbers')
        self.status.configure(text=f'Saved to {path}')
        self.app._open(path, 'revision checklist')

    def export_ecn(self):
        messagebox.showinfo('Portfolio demo', 'Company-specific change forms are not included in this demo.')

    def show_flat(self):
        number = self.number
        self.close()
        self.app.show_revision_list(number)
        self.app.lift()
        self.app.focus_force()

class ColumnsPanel(ctk.CTkFrame):

    def __init__(self, app):
        super().__init__(app, corner_radius=12, border_width=2, border_color=MODE_ACCENTS[app.mode]['normal'])
        self.app = app
        self.boxes = {}
        self.place(relx=0.98, rely=0.02, anchor='ne')
        self.grid_columnconfigure(0, weight=1)
        header = ctk.CTkFrame(self, fg_color='transparent')
        header.grid(row=0, column=0, sticky='ew', padx=16, pady=(14, 0))
        header.grid_columnconfigure(0, weight=1)
        what = 'parts' if app.mode == 'components' else 'BOMs'
        ctk.CTkLabel(header, text='Columns shown', font=ctk.CTkFont(family=FONT, size=15, weight='bold'), anchor='w').grid(row=0, column=0, sticky='w')
        ctk.CTkButton(header, text='Close', width=70, height=26, font=ctk.CTkFont(family=FONT, size=12), command=self.close).grid(row=0, column=1, sticky='e', padx=(12, 0))
        ctk.CTkLabel(header, text=f'in {what} results', font=ctk.CTkFont(family=FONT, size=12), text_color=('gray35', 'gray70'), anchor='w').grid(row=1, column=0, columnspan=2, sticky='w')
        shown = set(app.visible_columns())
        body = ctk.CTkFrame(self, fg_color='transparent')
        body.grid(row=1, column=0, sticky='ew', padx=16, pady=(8, 4))
        for (row, (key, heading, *_)) in enumerate(MODE_COLUMNS[app.mode]):
            variable = tk.BooleanVar(value=key in shown)
            box = ctk.CTkCheckBox(body, text=heading, variable=variable, width=210, checkbox_width=18, checkbox_height=18, font=ctk.CTkFont(family=FONT, size=12), command=self.apply)
            box.grid(row=row, column=0, sticky='w', pady=3)
            self.boxes[key] = variable
        footer = ctk.CTkFrame(self, fg_color='transparent')
        footer.grid(row=2, column=0, sticky='ew', padx=16, pady=(4, 14))
        ctk.CTkButton(footer, text='Show all', width=100, height=28, font=ctk.CTkFont(family=FONT, size=12), command=self.show_all).pack(side='left')
        ctk.CTkLabel(footer, text='  saved as you tick', font=ctk.CTkFont(family=FONT, size=11), text_color=('gray45', 'gray60')).pack(side='left')
        self.bind('<Escape>', lambda event: self.close())
        self.after(60, self.lift)

    def apply(self):
        keys = [key for (key, variable) in self.boxes.items() if variable.get()]
        if not keys:
            for (key, variable) in self.boxes.items():
                variable.set(key in self.app.visible_columns())
            return
        self.app.show_columns(keys)

    def show_all(self):
        for variable in self.boxes.values():
            variable.set(True)
        self.apply()

    def close(self):
        self.destroy()
        self.app.columns_panel = None

class NetworkCard(ctk.CTkFrame):
    WORDING = {'offline': ('Demo data is unavailable.', 'The local synthetic search index is missing. Regenerate it with tools/build_demo_data.py, then choose Retry.'), 'no index': ('The sample index is missing.', 'The local demo index was not found at:\n\n{where}\n\nRegenerate the synthetic dataset and restart the demo.')}

    def __init__(self, app, state, retry):
        super().__init__(app, corner_radius=12, border_width=2, border_color=MODE_ACCENTS[app.mode]['normal'])
        self.app = app
        (title, body) = self.WORDING[state]
        body = body.format(where=os.path.dirname(demo_config.paths()['cells']))
        self.place(relx=0.5, rely=0.5, anchor='center')
        self.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(self, text=title, anchor='w', font=ctk.CTkFont(family=FONT, size=17, weight='bold')).grid(row=0, column=0, sticky='w', padx=24, pady=(20, 6))
        ctk.CTkLabel(self, text=body, anchor='w', justify='left', wraplength=560, font=ctk.CTkFont(family=FONT, size=13), text_color=('gray35', 'gray70')).grid(row=1, column=0, sticky='w', padx=24)
        buttons = ctk.CTkFrame(self, fg_color='transparent')
        buttons.grid(row=2, column=0, sticky='e', padx=24, pady=(16, 20))
        ctk.CTkButton(buttons, text='Retry', width=110, height=32, font=ctk.CTkFont(family=FONT, size=12), command=lambda : (self.close(), retry())).pack(side='left', padx=(0, 8))
        ctk.CTkButton(buttons, text='Close', width=90, height=32, fg_color=('gray75', 'gray30'), hover_color=('gray68', 'gray38'), font=ctk.CTkFont(family=FONT, size=12), command=self.close).pack(side='left')
        self.bind('<Escape>', lambda event: self.close())
        self.after(60, self.lift)

    def close(self):
        self.destroy()
        self.app.network_card = None

def self_check():
    lines = [f'{VERSION} self-check, {datetime.now():%Y-%m-%d %H:%M:%S}']
    ok = True
    try:
        lines.append(f'profile        : {demo_config.active_name()}')
        config = demo_config.paths()
        lines.append(f"index folder   : {os.path.dirname(config['cells'])}")
        lines.append(f'network state  : {network_state(config)}')
        for (label, path) in (('activity log', demo_config.usage_log()), ('feedback', demo_config.feedback_file())):
            try:
                existed = os.path.exists(path)
                os.makedirs(os.path.dirname(path), exist_ok=True)
                with open(path, 'a', encoding='utf-8'):
                    pass
                state = 'exists and is writable' if existed else 'created just now, writable'
            except OSError as error:
                state = f'CANNOT BE WRITTEN - nothing is being saved ({error})'
            lines.append(f'{label:<15}: {path} - {state}')
        wanted = {'projects': config.get('projects'), 'project folders': config.get('project_folders'), 'obsolete BOMs': demo_config.obsolete_index()}
        (chosen, sampled) = ({}, set())
        for (table_name, table_path) in wanted.items():
            (found, is_sample) = demo_config.resolve_table(table_path)
            chosen[table_name] = found
            if is_sample:
                sampled.add(table_name)
        index = demo_index.load(config['cells'], config['files'], config['pdfs'], projects_path=chosen['projects'], project_folders_path=chosen['project folders'], obsolete_path=chosen['obsolete BOMs'], dfmea_path=config.get('dfmea'), prototype_dir=config.get('prototype_index'))
        index.sampled = sampled
        lines.append(f'workbooks      : {index.workbook_count:,}')
        lines.append(f'built at       : {index.built_at}')
        days = index.age_in_days()
        if days is None:
            verdict = 'no readable build date - no warning possible'
        elif stale_index_age(index) is not None:
            verdict = f'{days} days - the red stale-index strip WILL show'
        elif not demo_config.warn_when_index_is_old():
            verdict = f'{days} days - this profile never warns'
        else:
            verdict = f'{days} days - fresh; the strip would show at {demo_config.STALE_INDEX_DAYS}'
        lines.append(f'index age      : {verdict}')
        results = index.search('bracket')
        lines.append(f'search test    : bracket -> {len(results):,} BOMs')
        index.warm_revision_walk()
        lines.append(f'assembly trees : {len(index._part_uses()):,} part numbers')
        lines.append(f'customtkinter  : {ctk.__file__}')
        lines.append(f"ECN projects   : {len(index.projects):,} ({chosen['projects'] or wanted['projects'] or 'no path'})")
        lines.append(f"obsolete BOMs  : {len(index.obsolete):,} ({chosen['obsolete BOMs'] or wanted['obsolete BOMs']})")
        lines.append(f"project folders: {len(index.project_folders):,} ({chosen['project folders'] or wanted['project folders'] or 'no path'})")
        lines.append('risk records   : intentionally omitted from this synthetic demo')
        lines.append('prototype data : intentionally omitted from this synthetic demo')
        lines.append('samples in use : %s' % (', '.join(sorted(sampled)) if sampled else 'none - this demo uses synthetic tables only'))
        document = demo_config.help_document()
        lines.append('work instruction: %s' % (document or 'MISSING - Help will fall back to the About box'))
        theme = os.path.join(os.path.dirname(ctk.__file__), 'assets', 'themes')
        lines.append(f'themes present : {os.path.isdir(theme)}')
        if not os.path.isdir(theme):
            ok = False
    except Exception as error:
        ok = False
        lines.append(f'FAILED         : {type(error).__name__}: {error}')
    lines.append('result         : ' + ('OK' if ok else 'FAILED'))
    report = '\n'.join(lines) + '\n'
    exe = demo_config.running_exe()
    if exe:
        try:
            with open(os.path.join(os.path.dirname(exe), 'selfcheck.txt'), 'w', encoding='utf-8') as handle:
                handle.write(report)
        except OSError:
            pass
    print(report)
    return 0 if ok else 1
APP_MUTEX = 'BOMSearchDemoRunning'
_app_mutex = None

def hold_app_mutex():
    global _app_mutex
    try:
        import ctypes
        kernel32 = ctypes.windll.kernel32
        kernel32.CreateMutexW.restype = ctypes.c_void_p
        _app_mutex = kernel32.CreateMutexW(None, False, APP_MUTEX)
    except (AttributeError, OSError):
        _app_mutex = None
    return _app_mutex

def release_app_mutex():
    global _app_mutex
    if _app_mutex:
        import ctypes
        ctypes.windll.kernel32.CloseHandle(ctypes.c_void_p(_app_mutex))
    _app_mutex = None

def main():
    if '--self-check' in sys.argv or os.environ.get('DEMO_SELF_CHECK'):
        sys.exit(self_check())
    hold_app_mutex()
    App().mainloop()
if __name__ == '__main__':
    main()
